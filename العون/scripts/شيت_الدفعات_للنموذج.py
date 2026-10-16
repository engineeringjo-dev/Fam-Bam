# -*- coding: utf-8 -*-
"""إضافة شيت ثالث «الدفعات» على نموذج الورشات المعتمد (صاحب المحل 08/10)
الاستعمال: python3 scripts/شيت_الدفعات_للنموذج.py <النموذج.xlsx> <الناتج.xlsx>"""
import sys, copy, openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
R = wb['المرتجعات']
if 'الدفعات' in wb.sheetnames: del wb['الدفعات']
ws = wb.create_sheet('الدفعات', index=2)
GREEN, GREEN_L = '00375623', '00E2EFDA'
N1, N2 = 7, 306  # 300 دفعة
ws.sheet_view.rightToLeft = True
ws.freeze_panes = 'A7'
def st(dst, srcc, fill=None, fmt=None, color=None):
    dst.font = copy.copy(srcc.font); dst.border = copy.copy(srcc.border); dst.alignment = copy.copy(srcc.alignment)
    dst.fill = PatternFill('solid', fgColor=fill) if fill else copy.copy(srcc.fill)
    dst.number_format = fmt or srcc.number_format
    if color: f = copy.copy(srcc.font); f.color = color; dst.font = f
# العنوان واسم الورشة والعدّاد
ws['A1'] = 'دفعات الورشة — محلات العون لمواد البناء'; st(ws['A1'], R['A1'], GREEN); ws.merge_cells('A1:H1')
for c in 'AB': st(ws[c + '2'], R[c + '2'])
ws['A2'] = 'الورشة'; ws['B2'] = R['B2'].value; ws.merge_cells('B2:H2')
for c in 'BCDEFGH': st(ws[c + '2'], R['B2'])
ws['A4'] = 'عدد الدفعات'; ws['B4'] = '=COUNTA(D%d:D%d)' % (N1, N2)
ws['C4'] = 'المدقَّقة'; ws['D4'] = '=COUNTIFS(D%d:D%d,"<>",G%d:G%d,"<>")' % (N1, N2, N1, N2)
ws['E4'] = 'مجموع الدفعات'; ws['F4'] = '=SUM(D%d:D%d)' % (N1, N2)
ws['G4'] = '=IF(B4=0,"⛔ ما في دفعات مُدخلة",IF(D4<B4,"⛔ باقي "&(B4-D4)&" دفعة بدون تدقيق","✅ كل الدفعات مدقَّقة"))'
for c, s in zip('ABCDEFGH', ['A4', 'B4', 'C4', 'B4', 'A4', 'B4', 'F4', 'F4']): st(ws[c + '4'], R[s])
ws['F4'].number_format = '0.000'; ws.merge_cells('G4:H4')
# العناوين
H = ['#', 'التاريخ', 'نوع الدفعة', 'المبلغ', 'من', 'الى', 'المدقق', 'ملاحظات']
W = [5, 13, 16, 17, 26, 26, 18, 38]
for i, (h, w) in enumerate(zip(H, W), 1):
    c = ws.cell(6, i, h); st(c, R['A6'], '007030A0' if h == 'ملاحظات' else ('00C65911' if h == 'المبلغ' else GREEN))
    ws.column_dimensions[c.column_letter].width = w
for r in range(N1, N2 + 1):
    ws.cell(r, 1, '=IF(D%d="","",ROW()-6)' % r); st(ws.cell(r, 1), R['A7'])
    st(ws.cell(r, 2), R['F7']); st(ws.cell(r, 3), R['C7']); st(ws.cell(r, 4), R['E7'], fmt='0.000')
    for col in (5, 6): st(ws.cell(r, col), R['B7'])
    st(ws.cell(r, 7), R['H7']); st(ws.cell(r, 8), R['I7'])
    ws.row_dimensions[r].height = 21
for r, h in ((1, 30), (2, 26), (4, 26), (6, 28)): ws.row_dimensions[r].height = h
ws.cell(N2 + 2, 1, 'الدفعة بتنكتب يوم ما وصلت · «من» = مين دفع · «الى» = مين استلم (مثال: م.عمر) · أي ملاحظة بعمود «ملاحظات»'); ws.merge_cells('A%d:H%d' % (N2 + 2, N2 + 2))
st(ws.cell(N2 + 2, 1), R['A208'])
rng = lambda c: '%s%d:%s%d' % (c, N1, c, N2)
dv = DataValidation(type='list', formula1='"كاش نقدي,تحويل كليك,شيك"', allow_blank=True, showDropDown=False,
                    promptTitle='نوع الدفعة', prompt='اختار من القائمة: كاش نقدي · تحويل كليك · شيك',
                    errorTitle='نوع غلط', error='اختار من القائمة بس', showErrorMessage=True, showInputMessage=True)
dv.add(rng('C')); ws.add_data_validation(dv)
dv = DataValidation(type='decimal', operator='greaterThan', formula1='0', promptTitle='المبلغ', prompt='المبلغ بالدينار — رقم بس',
                    errorTitle='المبلغ', error='اكتب رقم أكبر من صفر', showErrorMessage=True, showInputMessage=True)
dv.add(rng('D')); ws.add_data_validation(dv)
dv = DataValidation(type='date', operator='between', formula1='DATE(2020,1,1)', formula2='DATE(2035,12,31)',
                    promptTitle='التاريخ: يوم/شهر/سنة', prompt='اكتب اليوم أول ثم الشهر ثم السنة\nمثال: 02/10/2026 = 2 تشرين أول',
                    errorTitle='تاريخ غلط', error='اكتب التاريخ يوم/شهر/سنة — مثال 02/10/2026', showErrorMessage=True, showInputMessage=True)
dv.add(rng('B')); ws.add_data_validation(dv)
dv = DataValidation(type='list', formula1='"عمر المصري,ابو علي"', allow_blank=True); dv.add(rng('G')); ws.add_data_validation(dv)
# تنبيه أحمر: دفعة بلا تاريخ/نوع/من/الى · أخضر: مدقَّقة
red = PatternFill('solid', fgColor='00FFC7CE'); ok = PatternFill('solid', fgColor=GREEN_L)
for c in 'BCEF':
    ws.conditional_formatting.add(rng(c), FormulaRule(formula=['AND($D%d<>"",%s%d="")' % (N1, c, N1)], fill=red))
ws.conditional_formatting.add('A%d:G%d' % (N1, N2), FormulaRule(formula=['$G%d<>""' % N1], fill=ok))
ws.conditional_formatting.add('G4', FormulaRule(formula=['LEFT(G4,1)="✅"'], fill=ok, font=Font(color='00375623', b=True)))
ws.conditional_formatting.add('G4', FormulaRule(formula=['LEFT(G4,1)="⛔"'], fill=red, font=Font(color='009C0006', b=True)))
ws.page_setup.orientation = 'landscape'; ws.print_title_rows = '1:6'; ws.print_area = 'A1:H%d' % (N2 + 2)
ws.sheet_properties.tabColor = '00375623'
# التعليمات
T = wb['التعليمات']; last = T.max_row
idx = next(r for r in range(1, last + 1) if T.cell(r, 2).value is None and (T.cell(r, 1).value or '').startswith('المدقِّق'))
moved = [str(m) for m in T.merged_cells.ranges if m.min_row >= idx]
for m in moved: T.unmerge_cells(m)
T.insert_rows(idx, 5)
for m in moved:
    c1, c2 = m.split(':'); T.merge_cells('%s%d:%s%d' % (c1[0], int(c1[1:]) + 5, c2[0], int(c2[1:]) + 5))
T.merge_cells('A%d:B%d' % (idx, idx))
T.cell(idx, 1, 'ورقة «الدفعات» — الفلوس اللي بتوصل (التبويب الأخضر)'); st(T.cell(idx, 1), T.cell(idx + 5, 1))
lines = ['كل دفعة بسطر: التاريخ · نوع الدفعة (كاش نقدي أو تحويل كليك أو شيك — من القائمة) · المبلغ.',
         '«من» = اسم اللي دفع · «الى» = اسم اللي استلم (مثال: م.عمر). الخانة بتصير حمرا إذا ناقصة.',
         'أي ملاحظة (مثل رقم الشيك أو تاريخ صرفه) بتنكتب بعمود «ملاحظات» البنفسجي.',
         'المدقق بيختار اسمه على كل دفعة بعد ما يتأكد منها — والعدّاد فوق بيصير أخضر لما كلها تتدقّق.']
for i, t in enumerate(lines):
    a, b = T.cell(idx + 1 + i, 1, 'د%d' % (i + 1)), T.cell(idx + 1 + i, 2, t)
    st(a, T['A4']); st(b, T['B4'])
for r in range(idx, idx + 5): T.row_dimensions[r].height = T.row_dimensions[4].height
for r in range(1, T.max_row + 1):
    v = T.cell(r, 2).value
    if isinstance(v, str) and 'لازم الاثنين يصيروا أخضر' in v: T.cell(r, 2).value = v.replace('لازم الاثنين يصيروا أخضر', 'لازم الثلاثة يصيروا أخضر')
wb.save(out); print('✓', out, wb.sheetnames)
