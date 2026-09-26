# -*- coding: utf-8 -*-
"""نظام الموظفين — محلات العون لمواد البناء (نسخة الساعات)
القواعد (26/09/2026 من صاحب المحل):
  • المدير: 125 دينار كل 10 أيام، بلا ساعات.
  • عمر المصري: السبت–الخميس 07:00–18:00 @1.250/ساعة · بعد 18:00 @1.500 · الجمعة 14:00–23:00 حسب الساعات @1.250.
  • عبدالعزيز: 1.000/ساعة بغض النظر عن اليوم · يقبض نهاية كل يوم.
  • تحميل الجبسمبورد: أي موظف يتحاسب على ساعاته فيه @2.500.
  • عمال المهام: بلا أسعار ثابتة — مقاولة متفق عليها لكل مهمة.
كل الأسعار بورقة «القواعد» ومربوطة بالمعادلات — غيّرها هناك وبس.
"""
import openpyxl, datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

OUT = '/home/user/Fam-Bam/العون/نظام_الموظفين.xlsx'
NAVY, TEAL, GOLD, PLUM, RUST = '1F4E6B', '0F6E6E', '8A6A00', '5B3A6B', '8B3A2F'
EDIT, AUTO, GREY = 'FFF6CC', 'EDEDED', 'F5F5F5'
LINE = Side(style='thin', color='9AA5AD'); THICK = Side(style='medium', color=NAVY)
BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
R0 = 6
LOG_ROWS, DAY_ROWS, PERIODS, TASK_ROWS = 300, 200, 12, 100
WEEK1 = dt.date(2026, 9, 26)          # السبت

def F(sz=11, b=False, c='000000'): return Font(name='Arial', size=sz, bold=b, color=c)
def C(h='center'): return Alignment(horizontal=h, vertical='center', wrap_text=True)
def fill(c): return PatternFill('solid', fgColor=c)
def col(i): return openpyxl.utils.get_column_letter(i)

wb = openpyxl.Workbook()

def sheet(title, banner, color, headers, widths, nrows, note=''):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True; ws.sheet_view.showGridLines = False
    n = len(headers); last = col(n)
    ws.merge_cells('A1:%s1' % last); ws['A1'] = banner
    ws['A1'].font = F(15, True, 'FFFFFF'); ws['A1'].alignment = C(); ws['A1'].fill = fill(color)
    ws.row_dimensions[1].height = 30
    ws.merge_cells('A2:%s2' % last); ws['A2'] = note
    ws['A2'].font = F(10, False, '444444'); ws['A2'].alignment = C('right'); ws.row_dimensions[2].height = 36
    ws.merge_cells('A3:%s3' % last)
    ws['A3'] = '🟨 صفراء = تتعبّى   ⬜ رمادية = تنحسب لحالها'
    ws['A3'].font = F(10, True, GOLD); ws['A3'].alignment = C('right')
    for k, w in enumerate(widths, start=1): ws.column_dimensions[col(k)].width = w
    for i, h in enumerate(headers, start=1):
        c = ws.cell(5, i, h); c.font = F(11, True, 'FFFFFF'); c.alignment = C(); c.fill = fill(color)
        c.border = Border(left=LINE, right=LINE, top=THICK, bottom=THICK)
    ws.row_dimensions[5].height = 32
    for r in range(R0, R0 + nrows):
        ws.row_dimensions[r].height = 20
        for cc in range(1, n + 1):
            c = ws.cell(r, cc); c.border = BOX; c.font = F(11); c.alignment = C(); c.fill = fill('FFFFFF')
    ws.freeze_panes = 'A6'
    return ws

def paint(ws, cols, lo, hi, color):
    for cl in cols:
        for r in range(lo, hi + 1): ws['%s%d' % (cl, r)].fill = fill(color)

def dv(ws, formula, rng):
    d = DataValidation(type='list', formula1=formula, allow_blank=True); ws.add_data_validation(d); d.add(rng)

def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)

# ═══ قائمة الأوقات: كل ربع ساعة بنظام 12 ساعة (تبدأ 6:00 AM لأنها بداية الدوام) ═══
def label(m):
    h, mi = divmod(m, 60)
    return '%d:%02d %s' % ((h % 12) or 12, mi, 'AM' if h < 12 else 'PM')
SLOTS = [(6 * 60 + 15 * k) % 1440 for k in range(96)]
TL = wb.create_sheet('الأوقات')
for i, m in enumerate(SLOTS, start=1):
    TL.cell(i, 1, label(m))
    TL.cell(i, 2, dt.time(m // 60, m % 60)).number_format = 'h:mm AM/PM'
TL.sheet_state = 'hidden'
TIME_LIST = "='الأوقات'!$A$1:$A$96"

def TV(ref):   # نص «7:15 PM» ← قيمة وقت
    return "INDEX('الأوقات'!$B$1:$B$96,MATCH(%s,'الأوقات'!$A$1:$A$96,0))" % ref

DAY = 'CHOOSE(WEEKDAY({d},1),"الأحد","الاثنين","الثلاثاء","الأربعاء","الخميس","الجمعة","السبت")'

# ═══════════════════ القواعد ═══════════════════
K = wb.create_sheet('القواعد')
K.sheet_view.rightToLeft = True; K.sheet_view.showGridLines = False
K.merge_cells('A1:D1'); K['A1'] = 'القواعد والأسعار — كل المعادلات بالملف بتقرأ من هون'
K['A1'].font = F(15, True, 'FFFFFF'); K['A1'].fill = fill(GOLD); K['A1'].alignment = C(); K.row_dimensions[1].height = 30
for k, w in zip('ABCD', [4, 46, 16, 52]): K.column_dimensions[k].width = w
for i, h in enumerate(['#', 'القاعدة', 'القيمة', 'الشرح'], start=1):
    c = K.cell(3, i, h); c.font = F(11, True, 'FFFFFF'); c.fill = fill(GOLD); c.alignment = C()
rules = [  # (اسم معرّف, القاعدة, القيمة, تنسيق, الشرح)
    ('MGR_SAL',   'المدير — راتب كل 10 أيام', 125, '#,##0.000', 'بلا ساعات عمل. كل 10 أيام دفعة.'),
    ('OMR_RATE',  'عمر المصري — أجر الساعة العادية', 1.25, '#,##0.000', 'السبت–الخميس من 07:00 لحد ساعة بداية الإضافي، والجمعة كل ساعاته.'),
    ('OMR_OT',    'عمر المصري — أجر الساعة بعد الدوام', 1.5, '#,##0.000', 'السبت–الخميس بعد الساعة المكتوبة تحت.'),
    ('OMR_OT_AT', 'عمر المصري — بداية الإضافي (الساعة)', dt.time(18, 0), 'h:mm AM/PM', 'أي ساعة بعدها بتنحسب @ الإضافي. الجمعة ما إلها إضافي.'),
    ('OMR_FRI_IN', 'عمر المصري — الجمعة من', dt.time(14, 0), 'h:mm AM/PM', 'دوام الجمعة المتفق عليه (للتنبيه بس، الحساب حسب الساعات الفعلية).'),
    ('OMR_FRI_OUT', 'عمر المصري — الجمعة إلى', dt.time(23, 0), 'h:mm AM/PM', ''),
    ('AZZ_RATE',  'عبدالعزيز — أجر الساعة', 1.0, '#,##0.000', 'بغض النظر عن اليوم أو الوقت. يقبض نهاية كل يوم.'),
    ('GYP_RATE',  'تحميل جبسمبورد — أجر الساعة', 2.5, '#,##0.000', 'أي موظف: ساعات التحميل بتنحسب @2.5 بدل أجره العادي.'),
    ('OMR_INST',  'عمر المصري — قسط أسبوعي ثابت (دين/بضاعة)', 0, '#,##0.000', 'ينخصم كل أسبوع من الصافي. 0 = ما عليه أقساط.'),
    ('AZZ_INST',  'عبدالعزيز — قسط أسبوعي ثابت', 0, '#,##0.000', ''),
    ('MGR_INST',  'المدير — قسط ثابت كل 10 أيام', 0, '#,##0.000', ''),
    ('ADV_CAP',   'سقف السلف — نسبة من مستحق الدورة', 0.5, '0%', 'تنبيه أحمر بالتسوية لو السلف تجاوزت. 0 = بلا سقف.'),
]
for i, (nm, rule, val, fmt, desc) in enumerate(rules, start=1):
    r = 3 + i
    K.cell(r, 1, i); K.cell(r, 2, rule); K.cell(r, 3, val); K.cell(r, 4, desc)
    for cc in range(1, 5):
        c = K.cell(r, cc); c.border = BOX; c.font = F(11); c.alignment = C('right' if cc in (2, 4) else 'center')
    K.cell(r, 3).fill = fill(EDIT); K.cell(r, 3).font = F(11, True); K.cell(r, 3).number_format = fmt
    K.cell(r, 4).font = F(10, False, '555555'); K.row_dimensions[r].height = 30
    name(nm, "'القواعد'!$C$%d" % r)
K['B17'] = 'عمال المهام: بلا أسعار ثابتة — كل مهمة مقاولة بسعرها بورقة «المهام».'
K['B17'].font = F(10, True, GOLD); K['B17'].alignment = C('right')
K['B18'] = 'الأسبوع بالملف = من السبت للجمعة. يوم القبض الأسبوعي: آخر يوم بالأسبوع (الجمعة) أو حسب ما تحدد.'
K['B18'].font = F(10, False, '555555'); K['B18'].alignment = C('right')

# ═══════════════════ ساعات عمر ═══════════════════
O = sheet('ساعات عمر', 'ساعات عمر المصري — سطر لكل يوم', NAVY,
          ['التاريخ', 'اليوم', 'دخول', 'خروج', 'منها تحميل جبسمبورد (ساعات)', 'إجمالي الساعات', 'ساعات عادية', 'ساعات إضافي', 'المستحق (دينار)', 'ملاحظات'],
          [13, 11, 12, 12, 15, 11, 11, 11, 13, 30], DAY_ROWS,
          'اكتب التاريخ واختر الدخول والخروج من القائمة (كل ربع ساعة · AM/PM). السبت–الخميس: الساعات بعد بداية الإضافي @الإضافي. الجمعة: كل الساعات @العادي. ساعات الجبسمبورد @2.5 بدل أجرها.')
paint(O, 'ACDEJ', R0, R0 + DAY_ROWS - 1, EDIT); paint(O, 'BFGHI', R0, R0 + DAY_ROWS - 1, AUTO)
def omar_formulas(ws, r):
    ws['B%d' % r] = '=IF(A{r}="","",%s)' % DAY.format(d='A{r}')
    ws['B%d' % r] = ws['B%d' % r].value.replace('{r}', str(r))
    ws['K%d' % r] = ('=IF(C{r}="","",%s)' % TV('C{r}')).replace('{r}', str(r))
    ws['L%d' % r] = ('=IF(D{r}="","",%s)' % TV('D{r}')).replace('{r}', str(r))
    ws['F%d' % r] = '=IF(OR(K{r}="",L{r}=""),"",ROUND(MOD(L{r}-K{r},1)*24,2))'.replace('{r}', str(r))
    # الجمعة: كل الساعات عادية · غيرها: العادي لحد OMR_OT_AT والباقي إضافي
    ws['G%d' % r] = ('=IF(F{r}="","",IF(WEEKDAY(A{r},1)=6,F{r},'
                     'ROUND(MAX(0,MIN(IF(L{r}<K{r},L{r}+1,L{r}),OMR_OT_AT)-K{r})*24,2)))').replace('{r}', str(r))
    ws['H%d' % r] = '=IF(F{r}="","",ROUND(F{r}-G{r},2))'.replace('{r}', str(r))
    # المستحق = عادي×سعر + إضافي×سعر + فرق الجبسمبورد (بياخذ من العادي أولاً ثم من الإضافي)
    ws['I%d' % r] = ('=IF(F{r}="","",ROUND(G{r}*OMR_RATE+H{r}*OMR_OT'
                     '+MIN(N(E{r}),G{r})*(GYP_RATE-OMR_RATE)+MAX(0,MIN(N(E{r}),F{r})-G{r})*(GYP_RATE-OMR_OT),3))').replace('{r}', str(r))
    for cl in 'KL': ws['%s%d' % (cl, r)].number_format = 'h:mm AM/PM'
    ws['A%d' % r].number_format = 'DD/MM/YYYY'; ws['I%d' % r].number_format = '#,##0.000'
    for cl in 'EFGH': ws['%s%d' % (cl, r)].number_format = '0.00'
for r in range(R0, R0 + DAY_ROWS): omar_formulas(O, r)
OE = R0 + DAY_ROWS - 1
# تنبيه: جمعة خارج 14:00–23:00 · جبسمبورد أكثر من الساعات
dv(O, TIME_LIST, 'C%d:D%d' % (R0, OE))
O.column_dimensions['K'].hidden = True; O.column_dimensions['L'].hidden = True
O.conditional_formatting.add('C%d:D%d' % (R0, OE), FormulaRule(formula=['AND($A%d<>"",$K%d<>"",$L%d<>"",WEEKDAY($A%d,1)=6,OR($K%d<OMR_FRI_IN,$L%d>OMR_FRI_OUT))' % (R0, R0, R0, R0, R0, R0)], fill=fill('FDE2C8')))
O.conditional_formatting.add('E%d:E%d' % (R0, OE), FormulaRule(formula=['AND($E%d<>"",$F%d<>"",$E%d>$F%d)' % (R0, R0, R0, R0)], fill=fill('FBD5D5')))
tot = OE + 1
O['H%d' % tot] = 'المجموع'; O['H%d' % tot].font = F(11, True); O['H%d' % tot].alignment = C()
O['I%d' % tot] = '=SUM(I%d:I%d)' % (R0, OE); O['I%d' % tot].number_format = '#,##0.000'; O['I%d' % tot].font = F(11, True); O['I%d' % tot].fill = fill(AUTO); O['I%d' % tot].border = BOX
ex = tot + 2
O['J%d' % ex] = 'مثال (مش بيانات حقيقية) — السبت 7:00 AM → 7:30 PM ومنها ساعتين جبسمبورد: 12.5 ساعة = 11 عادي + 1.5 إضافي → 18.500'
O['J%d' % ex].font = F(10, False, '666666'); O['J%d' % ex].alignment = C('right')
O['A%d' % ex] = WEEK1; O['C%d' % ex] = '7:00 AM'; O['D%d' % ex] = '7:30 PM'; O['E%d' % ex] = 2
omar_formulas(O, ex)
for cl in 'ABCDEFGHI': O['%s%d' % (cl, ex)].font = F(10, False, '888888'); O['%s%d' % (cl, ex)].border = BOX

# ═══════════════════ ساعات عبدالعزيز ═══════════════════
Z = sheet('ساعات عبدالعزيز', 'ساعات عبدالعزيز — سطر لكل يوم · يقبض نهاية اليوم', TEAL,
          ['التاريخ', 'اليوم', 'دخول', 'خروج', 'منها تحميل جبسمبورد (ساعات)', 'إجمالي الساعات', 'المستحق (دينار)', 'قبض؟', 'المقبوض (دينار)', 'ملاحظات'],
          [13, 11, 12, 12, 15, 11, 13, 9, 13, 30], DAY_ROWS,
          'دينار للساعة بغض النظر عن اليوم. ساعات الجبسمبورد @2.5. لما يقبض آخر اليوم اختر «نعم» واكتب المقبوض — الفرق بيظهر بالتسوية.')
paint(Z, 'ACDEHIJ', R0, R0 + DAY_ROWS - 1, EDIT); paint(Z, 'BFG', R0, R0 + DAY_ROWS - 1, AUTO)
def azz_formulas(ws, r):
    ws['B%d' % r] = ('=IF(A{r}="","",%s)' % DAY.format(d='A{r}')).replace('{r}', str(r))
    ws['K%d' % r] = ('=IF(C{r}="","",%s)' % TV('C{r}')).replace('{r}', str(r))
    ws['L%d' % r] = ('=IF(D{r}="","",%s)' % TV('D{r}')).replace('{r}', str(r))
    ws['F%d' % r] = '=IF(OR(K{r}="",L{r}=""),"",ROUND(MOD(L{r}-K{r},1)*24,2))'.replace('{r}', str(r))
    ws['G%d' % r] = '=IF(F{r}="","",ROUND(F{r}*AZZ_RATE+MIN(N(E{r}),F{r})*(GYP_RATE-AZZ_RATE),3))'.replace('{r}', str(r))
    for cl in 'KL': ws['%s%d' % (cl, r)].number_format = 'h:mm AM/PM'
    ws['A%d' % r].number_format = 'DD/MM/YYYY'
    for cl in 'GI': ws['%s%d' % (cl, r)].number_format = '#,##0.000'
    for cl in 'EF': ws['%s%d' % (cl, r)].number_format = '0.00'
for r in range(R0, R0 + DAY_ROWS): azz_formulas(Z, r)
ZE = R0 + DAY_ROWS - 1
dv(Z, '"نعم,لا"', 'H%d:H%d' % (R0, ZE))
dv(Z, TIME_LIST, 'C%d:D%d' % (R0, ZE))
Z.column_dimensions['K'].hidden = True; Z.column_dimensions['L'].hidden = True
Z.conditional_formatting.add('I%d:I%d' % (R0, ZE), FormulaRule(formula=['AND($H%d="نعم",$I%d<>$G%d)' % (R0, R0, R0)], fill=fill('FDE2C8')))
tot = ZE + 1
Z['F%d' % tot] = 'المجموع'; Z['F%d' % tot].font = F(11, True); Z['F%d' % tot].alignment = C()
for cl in 'GI':
    Z['%s%d' % (cl, tot)] = '=SUM(%s%d:%s%d)' % (cl, R0, cl, ZE); Z['%s%d' % (cl, tot)].number_format = '#,##0.000'
    Z['%s%d' % (cl, tot)].font = F(11, True); Z['%s%d' % (cl, tot)].fill = fill(AUTO); Z['%s%d' % (cl, tot)].border = BOX
ex = tot + 2
Z['J%d' % ex] = 'مثال — 8:00 AM → 4:00 PM ومنها ساعة جبسمبورد: 8 ساعات → 9.500 · قبض 9.500'
Z['J%d' % ex].font = F(10, False, '666666'); Z['J%d' % ex].alignment = C('right')
Z['A%d' % ex] = WEEK1; Z['C%d' % ex] = '8:00 AM'; Z['D%d' % ex] = '4:00 PM'; Z['E%d' % ex] = 1; Z['H%d' % ex] = 'نعم'; Z['I%d' % ex] = 9.5
azz_formulas(Z, ex)
for cl in 'ABCDEFGHI': Z['%s%d' % (cl, ex)].font = F(10, False, '888888'); Z['%s%d' % (cl, ex)].border = BOX

# ═══════════════════ السلف والديون ═══════════════════
L = sheet('السلف والديون', 'السلف والديون والبضاعة والدفعات — سطر لكل حركة', RUST,
          ['التاريخ', 'الموظف', 'النوع', 'المبلغ (دينار)', 'البيان'],
          [13, 18, 18, 14, 50], LOG_ROWS,
          'سلفة = كاش وسط الدورة بينخصم من الصافي · دين = مبلغ يُسدَّد بأقساط · بضاعة = بسعرها المتفق · دفع = ما قبضه راتباً أو صافي دورة · تسديد = دفع هو للمحل · رصيد سابق (له/عليه) = الأرصدة الافتتاحية.')
LE = R0 + LOG_ROWS - 1
paint(L, 'ABCDE', R0, LE, EDIT)
for r in range(R0, R0 + LOG_ROWS):
    L['A%d' % r].number_format = 'DD/MM/YYYY'; L['D%d' % r].number_format = '#,##0.000'
dv(L, '"عمر المصري,عبدالعزيز,المدير"', 'B%d:B%d' % (R0, LE))
dv(L, '"سلفة,دين,بضاعة,دفع,تسديد,رصيد سابق له,رصيد سابق عليه"', 'C%d:C%d' % (R0, LE))
ex = LE + 2
L['E%d' % ex] = 'مثال: 28/09/2026 · عمر المصري · سلفة · 20.000 · «كاش من الصندوق»'
L['E%d' % ex].font = F(10, False, '666666'); L['E%d' % ex].alignment = C('right')
name('LOG_D', "'السلف والديون'!$A$%d:$A$%d" % (R0, LE)); name('LOG_E', "'السلف والديون'!$B$%d:$B$%d" % (R0, LE))
name('LOG_T', "'السلف والديون'!$C$%d:$C$%d" % (R0, LE)); name('LOG_A', "'السلف والديون'!$D$%d:$D$%d" % (R0, LE))
name('OMR_D', "'ساعات عمر'!$A$%d:$A$%d" % (R0, OE)); name('OMR_P', "'ساعات عمر'!$I$%d:$I$%d" % (R0, OE))
name('AZZ_D', "'ساعات عبدالعزيز'!$A$%d:$A$%d" % (R0, ZE)); name('AZZ_P', "'ساعات عبدالعزيز'!$G$%d:$G$%d" % (R0, ZE))
name('AZZ_PAID', "'ساعات عبدالعزيز'!$I$%d:$I$%d" % (R0, ZE))

# ═══════════════════ المهام (مقاولة) ═══════════════════
T = sheet('المهام', 'عمال المهام — مقاولة لكل مهمة', PLUM,
          ['التاريخ', 'اسم العامل', 'المهمة', 'السعر المتفق (دينار)', 'دُفع؟', 'تاريخ الدفع', 'ملاحظات'],
          [13, 20, 40, 15, 9, 13, 30], TASK_ROWS,
          'بلا أسعار ثابتة: كل مهمة بسعرها المتفق مع العامل. المهمة اللي لسا ما اندفعت بتظهر بالتسوية كمستحق.')
TE = R0 + TASK_ROWS - 1
paint(T, 'ABCDEFG', R0, TE, EDIT)
for r in range(R0, R0 + TASK_ROWS):
    T['A%d' % r].number_format = 'DD/MM/YYYY'; T['F%d' % r].number_format = 'DD/MM/YYYY'; T['D%d' % r].number_format = '#,##0.000'
dv(T, '"نعم,لا"', 'E%d:E%d' % (R0, TE))
tot = TE + 1
T['C%d' % tot] = 'مجموع المهام غير المدفوعة'; T['C%d' % tot].font = F(11, True); T['C%d' % tot].alignment = C('right')
T['D%d' % tot] = '=SUMIFS(D%d:D%d,E%d:E%d,"لا")' % (R0, TE, R0, TE); T['D%d' % tot].number_format = '#,##0.000'
T['D%d' % tot].font = F(11, True); T['D%d' % tot].fill = fill(AUTO); T['D%d' % tot].border = BOX
T['G%d' % (tot + 2)] = 'مثال: 27/09/2026 · محمد · تنزيل طبلية جبسمبورد · 10.000 · نعم · 27/09/2026'
T['G%d' % (tot + 2)].font = F(10, False, '666666'); T['G%d' % (tot + 2)].alignment = C('right')

# ═══════════════════ المدير ═══════════════════
M = sheet('المدير', 'المدير — 125 دينار كل 10 أيام', GOLD,
          ['#', 'من', 'إلى', 'الراتب', 'سلف الفترة', 'قسط', 'الصافي', 'دُفع؟', 'المدفوع', 'ملاحظات'],
          [5, 13, 13, 12, 12, 10, 12, 9, 12, 28], PERIODS,
          'كل سطر فترة 10 أيام. أول تاريخ بس بتكتبه، والباقي بيتسلسل لحاله. السلف بتنجمع من ورقة «السلف والديون» حسب تواريخ الفترة.')
ME = R0 + PERIODS - 1
paint(M, 'HIJ', R0, ME, EDIT); paint(M, 'ACDEFG', R0, ME, AUTO); M['B%d' % R0].fill = fill(EDIT)
M['B%d' % R0] = WEEK1
for r in range(R0, ME + 1):
    M['A%d' % r] = '=IF(B%d="","",ROW()-%d)' % (r, R0 - 1)
    if r > R0: M['B%d' % r] = '=IF(C%d="","",C%d+1)' % (r - 1, r - 1)
    M['C%d' % r] = '=IF(B%d="","",B%d+9)' % (r, r)
    M['D%d' % r] = '=IF(B%d="","",MGR_SAL)' % r
    M['E%d' % r] = '=IF(B%d="","",SUMIFS(LOG_A,LOG_E,"المدير",LOG_T,"سلفة",LOG_D,">="&B%d,LOG_D,"<="&C%d))' % (r, r, r)
    M['F%d' % r] = '=IF(B%d="","",MGR_INST)' % r
    M['G%d' % r] = '=IF(B%d="","",D%d-E%d-F%d)' % (r, r, r, r)
    for cl in 'BC': M['%s%d' % (cl, r)].number_format = 'DD/MM/YYYY'
    for cl in 'DEFGI': M['%s%d' % (cl, r)].number_format = '#,##0.000'
dv(M, '"نعم,لا"', 'H%d:H%d' % (R0, ME))
M.conditional_formatting.add('G%d:G%d' % (R0, ME), FormulaRule(formula=['AND($G%d<>"",$G%d<0)' % (R0, R0)], font=Font(color='C00000', bold=True)))
M['J%d' % (ME + 2)] = 'لما تدفع الفترة: «نعم» + المبلغ هون، وسجّلها كمان «دفع» بورقة السلف والديون عشان الكشف.'
M['J%d' % (ME + 2)].font = F(10, False, '666666'); M['J%d' % (ME + 2)].alignment = C('right')

# ═══════════════════ التسوية الأسبوعية ═══════════════════
S = wb.create_sheet('التسوية الأسبوعية', 0)
S.sheet_view.rightToLeft = True; S.sheet_view.showGridLines = False
S.merge_cells('A1:I1'); S['A1'] = 'التسوية الأسبوعية — من السبت للجمعة'
S['A1'].font = F(15, True, 'FFFFFF'); S['A1'].fill = fill(NAVY); S['A1'].alignment = C(); S.row_dimensions[1].height = 30
for k, w in zip('ABCDEFGHI', [4, 18, 12, 14, 12, 12, 14, 14, 30]): S.column_dimensions[k].width = w
S['B3'] = 'بداية الأسبوع (السبت)'; S['B3'].font = F(11, True); S['B3'].alignment = C('right')
S['C3'] = WEEK1; S['C3'].number_format = 'DD/MM/YYYY'; S['C3'].fill = fill(EDIT); S['C3'].font = F(12, True); S['C3'].border = BOX
S['D3'] = 'نهايته'; S['D3'].font = F(11, True)
S['E3'] = '=C3+6'; S['E3'].number_format = 'DD/MM/YYYY'; S['E3'].fill = fill(AUTO); S['E3'].border = BOX
S['F3'] = '=IF(WEEKDAY(C3,1)<>7,"⚠️ التاريخ مش سبت","")'; S['F3'].font = F(10, True, 'C00000')
hdr = ['#', 'الموظف', 'الساعات', 'المستحق', 'السلف', 'الأقساط', 'مقبوض يومياً', 'الصافي للدفع', 'ملاحظة']
for i, h in enumerate(hdr, start=1):
    c = S.cell(5, i, h); c.font = F(11, True, 'FFFFFF'); c.fill = fill(NAVY); c.alignment = C()
    c.border = Border(left=LINE, right=LINE, top=THICK, bottom=THICK)
S.row_dimensions[5].height = 30
rows = [('عمر المصري', 'OMR_D', 'OMR_P', None, 'OMR_INST', "'ساعات عمر'!$F$%d:$F$%d" % (R0, OE)),
        ('عبدالعزيز', 'AZZ_D', 'AZZ_P', 'AZZ_PAID', 'AZZ_INST', "'ساعات عبدالعزيز'!$F$%d:$F$%d" % (R0, ZE))]
for i, (emp, dn, pn, paidn, inst, hrs) in enumerate(rows):
    r = 6 + i
    S['A%d' % r] = i + 1; S['B%d' % r] = emp
    S['C%d' % r] = '=SUMIFS(%s,%s,">="&$C$3,%s,"<="&$E$3)' % (hrs, dn, dn)
    S['D%d' % r] = '=SUMIFS(%s,%s,">="&$C$3,%s,"<="&$E$3)' % (pn, dn, dn)
    S['E%d' % r] = '=SUMIFS(LOG_A,LOG_E,B%d,LOG_T,"سلفة",LOG_D,">="&$C$3,LOG_D,"<="&$E$3)' % r
    S['F%d' % r] = '=IF(D%d=0,0,%s)' % (r, inst)
    S['G%d' % r] = '=SUMIFS(%s,%s,">="&$C$3,%s,"<="&$E$3)' % (paidn, dn, dn) if paidn else 0
    S['H%d' % r] = '=D%d-E%d-F%d-G%d' % (r, r, r, r)
    S['I%d' % r] = '=IF(AND(ADV_CAP>0,D%d>0,E%d>ADV_CAP*D%d),"⚠️ السلف تجاوزت السقف","")' % (r, r, r)
    for cc in range(1, 10):
        c = S.cell(r, cc); c.border = BOX; c.font = F(11); c.alignment = C(); c.fill = fill(AUTO)
    S['C%d' % r].number_format = '0.00'
    for cl in 'DEFGH': S['%s%d' % (cl, r)].number_format = '#,##0.000'
    S['H%d' % r].font = F(12, True); S['I%d' % r].font = F(10, True, 'C00000'); S['I%d' % r].alignment = C('right')
S.conditional_formatting.add('H6:H7', FormulaRule(formula=['$H6<0'], font=Font(color='C00000', bold=True)))
S['B9'] = 'المدير: بيتحاسب كل 10 أيام بورقة «المدير» — مش أسبوعي.'; S['B9'].font = F(10, False, '555555'); S['B9'].alignment = C('right')
S['B10'] = 'عمال المهام: المهام غير المدفوعة بورقة «المهام»:'; S['B10'].font = F(10, False, '555555'); S['B10'].alignment = C('right')
S['E10'] = "='المهام'!D%d" % (TE + 1); S['E10'].number_format = '#,##0.000'; S['E10'].fill = fill(AUTO); S['E10'].border = BOX; S['E10'].font = F(11, True)
# كشف الرصيد الجاري لكل موظف
S['B12'] = 'الرصيد الجاري لكل موظف (من أول الملف لليوم)'; S['B12'].font = F(12, True, TEAL)
hdr2 = ['#', 'الموظف', 'إجمالي المستحق', 'رصيد سابق له', 'تسديدات', 'سلف + ديون + بضاعة', 'مدفوع له', 'الرصيد', 'معناه']
for i, h in enumerate(hdr2, start=1):
    c = S.cell(13, i, h); c.font = F(10, True, 'FFFFFF'); c.fill = fill(TEAL); c.alignment = C(); c.border = BOX
S.row_dimensions[13].height = 30
bal = [('عمر المصري', '=SUM(OMR_P)'), ('عبدالعزيز', '=SUM(AZZ_P)'),
       ('المدير', "=SUMIFS('المدير'!$D$%d:$D$%d,'المدير'!$C$%d:$C$%d,\"<=\"&TODAY())" % (R0, ME, R0, ME))]
for i, (emp, earned) in enumerate(bal):
    r = 14 + i
    S['A%d' % r] = i + 1; S['B%d' % r] = emp; S['C%d' % r] = earned
    S['D%d' % r] = '=SUMIFS(LOG_A,LOG_E,B%d,LOG_T,"رصيد سابق له")' % r
    S['E%d' % r] = '=SUMIFS(LOG_A,LOG_E,B%d,LOG_T,"تسديد")' % r
    S['F%d' % r] = ('=SUMIFS(LOG_A,LOG_E,B{r},LOG_T,"سلفة")+SUMIFS(LOG_A,LOG_E,B{r},LOG_T,"دين")'
                    '+SUMIFS(LOG_A,LOG_E,B{r},LOG_T,"بضاعة")+SUMIFS(LOG_A,LOG_E,B{r},LOG_T,"رصيد سابق عليه")').replace('{r}', str(r))
    paid_daily = '+SUM(AZZ_PAID)' if emp == 'عبدالعزيز' else ''
    S['G%d' % r] = '=SUMIFS(LOG_A,LOG_E,B%d,LOG_T,"دفع")%s' % (r, paid_daily)
    S['H%d' % r] = '=C%d+D%d+E%d-F%d-G%d' % (r, r, r, r, r)
    S['I%d' % r] = '=IF(H%d>0.0005,"المحل مدين إله",IF(H%d<-0.0005,"هو مدين للمحل","متساوي"))' % (r, r)
    for cc in range(1, 10):
        c = S.cell(r, cc); c.border = BOX; c.font = F(11); c.alignment = C(); c.fill = fill(AUTO)
    for cl in 'CDEFGH': S['%s%d' % (cl, r)].number_format = '#,##0.000;[Red]-#,##0.000;-'
    S['H%d' % r].font = F(12, True)
S['B18'] = 'الرصيد = المستحق + رصيد سابق له + تسديداته − (سلف + ديون + بضاعة + رصيد سابق عليه) − المدفوع له. المدفوع لعبدالعزيز = المقبوض اليومي من ورقته + أي «دفع» بالسجل.'
S.merge_cells('B18:I18'); S['B18'].font = F(9, False, '555555'); S['B18'].alignment = C('right'); S.row_dimensions[18].height = 30

# ═══════════════════ التعليمات ═══════════════════
I = wb.create_sheet('التعليمات')
I.sheet_view.rightToLeft = True; I.sheet_view.showGridLines = False
I.column_dimensions['A'].width = 5; I.column_dimensions['B'].width = 115
I['B1'] = 'كيف يشتغل الملف'; I['B1'].font = F(15, True, NAVY)
steps = [
    ('القواعد', None),
    ('1', 'كل الأسعار والساعات بورقة «القواعد». غيّر أي رقم هناك وبيتغيّر الحساب بكل الملف. ما تكتب أسعار بالأوراق الثانية.'),
    ('كل يوم', None),
    ('2', '«ساعات عمر»: التاريخ، واختر الدخول والخروج من القائمة المنسدلة (كل ربع ساعة، AM/PM). السبت–الخميس: اللي بعد 18:00 بينحسب إضافي @1.5. الجمعة: كل الساعات @1.25 (ولو دوامه برّا 14:00–23:00 بيتلوّن برتقالي للتنبيه).'),
    ('3', '«ساعات عبدالعزيز»: نفس الشي @1.0، وآخر اليوم لما يقبض: «نعم» + المبلغ.'),
    ('4', 'ساعات تحميل الجبسمبورد: اكتبها بعمودها (هي جزء من ساعات اليوم، مش زيادة عليها) وبتنحسب @2.5 بدل الأجر العادي.'),
    ('5', 'أي سلفة أو دين أو بضاعة أو دفعة: سطر بورقة «السلف والديون». الأرصدة القائمة اليوم: نوعها «رصيد سابق له» أو «رصيد سابق عليه».'),
    ('6', 'عمال المهام: سطر لكل مهمة بسعرها المتفق، و«دُفع؟».'),
    ('آخر الأسبوع', None),
    ('7', '«التسوية الأسبوعية»: اكتب تاريخ السبت وبيطلعلك لكل واحد الساعات والمستحق والسلف والأقساط والصافي للدفع. عبدالعزيز بيطلع الفرق بين مستحقه ومقبوضه اليومي.'),
    ('8', 'لما تدفع الصافي: سجّله «دفع» بورقة السلف والديون عشان الرصيد الجاري يضبط.'),
    ('9', 'المدير: ورقة «المدير» كل 10 أيام، 125 دينار ناقص سلف الفترة والقسط.'),
    ('على أودو', None),
    ('10', 'هذا الملف هو ورقة العمل اليومية. مدير الموظفين بينقل منه لأودو: استحقاق كل دورة + السلف + الدفعات، على الحساب الجاري لكل موظف — بعد ما تشوف المسودة.'),
]
r = 3
for a, b in steps:
    if b is None: I['B%d' % r] = a; I['B%d' % r].font = F(12, True, TEAL); r += 1; continue
    I['A%d' % r] = a; I['A%d' % r].font = F(11, True); I['A%d' % r].alignment = C()
    I['B%d' % r] = b; I['B%d' % r].font = F(11); I['B%d' % r].alignment = C('right'); I.row_dimensions[r].height = 34; r += 1

del wb['Sheet']
wb.save(OUT); print('saved', OUT)
