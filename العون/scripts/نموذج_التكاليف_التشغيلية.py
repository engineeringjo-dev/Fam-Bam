# -*- coding: utf-8 -*-
"""نموذج «التكاليف التشغيلية» — محلات العون (صاحب المحل 09/10 و 10/10)
🔴 شيت واحد «التكاليف» (10/10 — «شيت وحده المواد والأيدي العاملة · الورشة تحتها موادها وأجورها»):
  الورشة (شغل داخلي) · البند (الأجرة سطر زي أي بند: «اسم العامل — الشغلة») · الكمية (رقم صحيح) · الإفرادي · الإجمالي (معادلة) ·
  التاريخ · اليوم (تلقائي) · الغرض (قائمة أو كتابة حرّة = أخرى · تلقائي: ورشة «ستاندات» ← أثاث وستاندات · الأجرة: «أجور عمال») ·
  أصل؟ (مصروف/أصل/زكاة وصدقة — تلقائي) · المدقق · ملاحظات إضافية.
  الورشة: مجموعها (مواد + أجور) 20 وفوق (حد الأصل 20 — قراره 10/10) ← كل أسطرها أصل · الشهر حسب التاريخ (عمال المهام كمان — بأي يوم حتى الجمعة).
شيت «الملخص» شهر بشهر + حسب الغرض + الورشات · «القوائم» مخفي (الغرض ← حساب أودو) · «التعليمات».
🔒 كل الخلايا مقفولة إلا خانات التعبئة · الملخص والتعليمات مقفولين كامل · كلمة السر **** · بيقراه `ترحيل_التكاليف.py`.
python3 scripts/نموذج_التكاليف_التشغيلية.py [الناتج.xlsx]"""
import sys, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.properties import CalcProperties

OUT = sys.argv[1] if len(sys.argv) > 1 else 'drive/نموذج_التكاليف_التشغيلية.xlsx'
PWD = __import__('os').environ.get('ALAWN_XLSX_PWD', ''); ASSET = 20; N = 800
FONT = 'Arial'
NAVY, ORG, GRN, PUR, BRN = '1F4E6B', 'C65911', '2E7D32', '7030A0', '8E5A2B'
YEL, PEACH, LIL, GRY, RED, OK, WS_FILL, WK_FILL = 'FFF2CC', 'FCE4D6', 'E4DFEC', 'F2F2F2', 'FFC7CE', 'E2EFDA', 'F8EFE6', 'EAF2FB'
DAYS = '"الأحد","الاثنين","الثلاثاء","الأربعاء","الخميس","الجمعة","السبت"'
MONTHS = ['كانون الثاني', 'شباط', 'آذار', 'نيسان', 'أيار', 'حزيران', 'تموز', 'آب', 'أيلول', 'تشرين الأول', 'تشرين الثاني', 'كانون الأول']
thin = Side(style='thin', color='A6A6A6'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda c: PatternFill('solid', fgColor=c)
UNLOCK, LOCK = Protection(locked=False), Protection(locked=True)

# الغرض ← حساب أودو (الغرض · حساب المصروف · حساب الأصل إذا 20+ · الفئة · أمثلة) — قائمة صاحب المحل 10/10
GROUPS = [
    ('ضيافة', '501106', '', 'مصروف', 'شاي · قهوة · مي · سكر · غدا للعمال'),
    ('هدايا', '500701', '', 'مصروف', 'هدية لزبون · عيّنات · كاتالوجات'),
    ('إيجار', '500402', '', 'مصروف', 'إيجار المستودعات — يوم الدفع بالمبلغ كامل (التحويش اليومي ما بينكتب)'),
    ('فاتورة كهرباء', '500403', '', 'مصروف', 'فاتورة الكهربا (والمي) — اكتب أي ساعة بالبند'),
    ('أثاث وستاندات', '501106', '100604', 'مصروف', 'ستاند · رف · مكتب · كرسي · كمبيوتر (%d وفوق = أصل)' % ASSET),
    ('صيانة المحل', '501103', '', 'مصروف', 'تصليح · سباكة · كهربا · دهان · لمبات · قطع غيار'),
    ('أجور عمال', '500312', '', 'مصروف', 'تعزيل · تنظيف · تحميل · نجّار · كهربجي من برّا — اسم العامل والشغلة بالبند'),
    ('أجور توصيل', '500802', '', 'مصروف', 'أجرة سيارة · توصيل · بنزين للتوصيل'),
    ('صدقة', '500910', '', 'زكاة وصدقة', 'صدقة · تبرع لجامع أو جمعية (كاش أو بضاعة بالكلفة)'),
    ('زكاة', '500912', '', 'زكاة وصدقة', 'زكاة المحل السنوية'),
    ('رسوم تراخيص', '500503', '', 'مصروف', 'رخصة المهن · البلدية · الغرفة'),
    ('توالف بضاعة', '500905', '', 'مصروف', 'بضاعة انكسرت أو خربت بالمحل (بالكلفة)'),
    ('أخرى', '500907', '', 'مصروف', 'أي إشي ثاني — أو اكتب الغرض بإيدك بالخانة (مثلاً «أكياس للمحل»)'),
]
NG = len(GROUPS); GMAP = "'القوائم'!$A$2:$D$%d" % (NG + 1)
wb = openpyxl.Workbook()


def cell(ws, r, c, v=None, fmt=None, fill=None, bold=False, right=False, edit=False, size=11, color='000000', wrap=False):
    x = ws.cell(r, c, v); x.border = B
    x.alignment = Alignment(horizontal='right' if right else 'center', vertical='center', wrap_text=wrap)
    if fmt: x.number_format = fmt
    if fill: x.fill = F(fill)
    x.font = Font(name=FONT, b=bold, size=size, color=color)
    x.protection = UNLOCK if edit else LOCK
    return x


def title(ws, text, rng, color, size=16, h=34):
    x = ws[rng.split(':')[0]]; x.value = text; x.font = Font(name=FONT, b=True, size=size, color='FFFFFF'); x.fill = F(color)
    x.alignment = Alignment(horizontal='center', vertical='center'); ws.merge_cells(rng); ws.row_dimensions[x.row].height = h


def header(ws, r, c, text, color, width=None, fc='FFFFFF', size=11):
    x = ws.cell(r, c, text); x.fill = F(color); x.font = Font(name=FONT, b=True, color=fc, size=size); x.border = B
    x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if width: ws.column_dimensions[x.column_letter].width = width


def dv(ws, typ, rng, f1=None, f2=None, title=None, msg=None, err=None, op=None, style='stop'):
    d = DataValidation(type=typ, formula1=f1, formula2=f2, operator=op, allow_blank=True, showDropDown=False, errorStyle=style,
                       promptTitle=title, prompt=msg, showInputMessage=bool(msg), errorTitle='تنبيه', error=err, showErrorMessage=bool(err))
    d.add(rng); ws.add_data_validation(d)


def protect(ws, full=False):
    ws.protection.sheet = True; ws.protection.set_password(PWD)
    ws.protection.formatColumns = ws.protection.formatRows = False   # توسيع الأعمدة والأسطر مسموح
    if full: ws.protection.selectLockedCells = True                  # الملخص والتعليمات: ولا خانة بتنحدد


# ═══ 1) التكاليف — شيت واحد للمواد والأجور (الأجرة سطر زي أي بند: «اسم العامل — الشغلة» · غرضها «أجور عمال») ═══
# ترتيب الأعمدة (10/10): A # · B التاريخ · C اليوم · D الورشة (شغل داخلي) · E البند · F الكمية · G الإفرادي · H الإجمالي ·
# I الغرض · J أصل؟ · K المدقق · L ملاحظات إضافية · مخفي: N الشهر · O السنة · P ترتيب الورشة · Q مادة/أجرة
M = wb.active; M.title = 'التكاليف'; M.sheet_view.rightToLeft = True; M.sheet_properties.tabColor = NAVY
cols = [('#', 6), ('التاريخ', 13), ('اليوم', 11), ('الورشة (شغل داخلي)', 22), ('البند', 36), ('الكمية', 9), ('الإفرادي', 11), ('الإجمالي', 12),
        ('الغرض', 19), ('أصل؟', 13), ('المدقق', 14), ('ملاحظات إضافية', 40)]
title(M, 'التكاليف التشغيلية — مواد وأجور عمال', 'A1:L1', NAVY)
HC = {'الإجمالي': ORG, 'ملاحظات إضافية': PUR, 'الورشة (شغل داخلي)': BRN}
for i, (h, w) in enumerate(cols, 1):
    header(M, 6, i, h, HC.get(h, NAVY), w)
M.row_dimensions[6].height = 32; M.freeze_panes = 'A7'
cell(M, 2, 1, 'الجهة', fill='DDE5EC', bold=True); cell(M, 2, 2, 'محلات العون — مصاريف المحل', fill=GRY, bold=True, right=True); M.merge_cells('B2:L2')
e1, e2 = 7, 6 + N
R = lambda c: '$%s$%d:$%s$%d' % (c, e1, c, e2)
WTOT = 'SUMIF(%s,{w},%s)' % (R('D'), R('H'))                     # مجموع الورشة (مواد + أجور)
BADWS = 'OR(ISNUMBER(SEARCH("صيانة",D{0})),ISNUMBER(SEARCH("تعزيل",D{0})),ISNUMBER(SEARCH("تنظيف",D{0})),ISNUMBER(SEARCH("تصليح",D{0})))'
for c, v, fm in ((1, 'البنود', None), (2, '=COUNTA(E%d:E%d)' % (e1, e2), '0'), (3, 'المجموع', None), (4, '=SUM(H%d:H%d)' % (e1, e2), '#,##0.00'),
                 (5, 'منها أصول', None), (6, '=SUMIF(J%d:J%d,"أصل",H%d:H%d)' % (e1, e2, e1, e2), '#,##0.00'),
                 (7, 'المدقَّقة', None), (8, '=SUMPRODUCT((E%d:E%d<>"")*(K%d:K%d<>""))' % (e1, e2, e1, e2), '0'),
                 (9, 'ناقص ⚠', None), (10, '=SUMPRODUCT((E%d:E%d<>"")*(LEFT(J%d:J%d,1)="⚠"))' % (e1, e2, e1, e2), '0')):
    cell(M, 4, c, v, fm, 'FFFFFF' if c % 2 == 0 else 'DDE5EC', True)
M.row_dimensions[4].height = 22
M.conditional_formatting.add('I4:J4', FormulaRule(formula=['$J$4>0'], fill=F(RED), font=Font(b=True, color='9C0006')))
for r in range(e1, e2 + 1):
    cell(M, r, 1, '=IF(E{0}="","",ROW()-6)'.format(r), fill=GRY)
    cell(M, r, 2, fmt='dd/mm/yyyy', edit=True)
    cell(M, r, 3, '=IF(B{0}="","",CHOOSE(WEEKDAY(B{0},1),{1}))'.format(r, DAYS), fill=GRY)
    cell(M, r, 4, fill=WS_FILL, bold=True, right=True, edit=True, wrap=True)
    cell(M, r, 5, right=True, edit=True, wrap=True)
    cell(M, r, 6, fmt='0', edit=True); cell(M, r, 7, fmt='#,##0.00', fill=YEL, edit=True)
    cell(M, r, 8, '=IF(OR(F{0}="",G{0}=""),"",ROUND(F{0}*G{0},2))'.format(r), '#,##0.00', PEACH, bold=True)
    cell(M, r, 9, '=IF(ISNUMBER(SEARCH("ستاند",D{0})),"أثاث وستاندات","")'.format(r), bold=True, edit=True, wrap=True)
    cell(M, r, 10, ('=IF(E{0}="","",IF(I{0}="","⚠ عبّي الغرض",IF(AND(D{0}<>"",{4}),"⚠ الصيانة بلا ورشة",IF(H{0}="","⚠ ناقص",'
                    'IF(D{0}<>"",IF({3}>={2},"أصل","مصروف"),'
                    'IFERROR(IF(VLOOKUP(I{0},{1},4,0)="زكاة وصدقة","زكاة وصدقة",IF(AND(VLOOKUP(I{0},{1},3,0)<>"",H{0}>={2}),"أصل","مصروف")),"مصروف"))))))'
                    ).format(r, GMAP, ASSET, WTOT.format(w='D%d' % r), BADWS.format(r)), fill=GRY, bold=True)
    cell(M, r, 11, edit=True); cell(M, r, 12, fill=LIL, right=True, edit=True, wrap=True)
    M.cell(r, 14, '=IF(B{0}="","",MONTH(B{0}))'.format(r)); M.cell(r, 15, '=IF(B{0}="","",YEAR(B{0}))'.format(r))
    M.cell(r, 16, '=IF(AND(D{0}<>"",COUNTIF(D${1}:D{0},D{0})=1),MAX(P${2}:P{3})+1,"")'.format(r, e1, e1 - 1, r - 1))
    M.cell(r, 17, '=IF(E{0}="","",IF(I{0}="أجور عمال","أجرة","مادة"))'.format(r))      # Q: مادة/أجرة (للملخص)
for c in 'NOPQ': M.column_dimensions[c].hidden = True
dv(M, 'list', 'I%d:I%d' % (e1, e2), "'القوائم'!$A$2:$A$%d" % (NG + 1), err='مش من القائمة — بينحسب «أخرى». تمام؟', style='information')
dv(M, 'list', 'K%d:K%d' % (e1, e2), '"عمر المصري,ابو علي,مهندس علي"')
dv(M, 'decimal', 'G%d:G%d' % (e1, e2), '0', err='رقم بس', op='greaterThanOrEqual')
dv(M, 'whole', 'F%d:F%d' % (e1, e2), '0', err='الكمية رقم صحيح بلا فاصلة (مثلاً 3) · الأجرة: 1', op='greaterThan')
dv(M, 'date', 'B%d:B%d' % (e1, e2), 'DATE(2020,1,1)', 'DATE(2035,12,31)', err='اكتب التاريخ يوم/شهر/سنة (مثلاً 07/10/2026)', op='between')
for c in 'BFG':
    M.conditional_formatting.add('%s%d:%s%d' % (c, e1, c, e2), FormulaRule(formula=['AND($E%d<>"",%s%d="")' % (e1, c, e1)], fill=F(RED)))
M.conditional_formatting.add('I%d:I%d' % (e1, e2), FormulaRule(formula=['AND($E%d<>"",I%d="")' % (e1, e1)], fill=F(RED)))
M.conditional_formatting.add('D%d:D%d' % (e1, e2), FormulaRule(formula=['AND(D{0}<>"",{1})'.format(e1, BADWS.format(e1))], fill=F(RED), font=Font(b=True, color='9C0006')))
M.conditional_formatting.add('J%d:J%d' % (e1, e2), FormulaRule(formula=['LEFT(J%d,1)="⚠"' % e1], fill=F(RED), font=Font(b=True, color='9C0006')))
M.conditional_formatting.add('L%d:L%d' % (e1, e2), FormulaRule(formula=['AND($I%d="أخرى",$L%d="")' % (e1, e1)], fill=F(RED)))
M.conditional_formatting.add('J%d:J%d' % (e1, e2), FormulaRule(formula=['J%d="أصل"' % e1], fill=F('DDEBF7'), font=Font(b=True, color='1F4E79')))
M.conditional_formatting.add('J%d:J%d' % (e1, e2), FormulaRule(formula=['J%d="زكاة وصدقة"' % e1], fill=F('FFF2CC'), font=Font(b=True, color='7F6000')))
M.conditional_formatting.add('A%d:K%d' % (e1, e2), FormulaRule(formula=['$K%d<>""' % e1], fill=F(OK)))
M.auto_filter.ref = 'A6:L%d' % e2                                   # فلتر: اختار ورشة أو غرض وبيطلعوا أسطره لحاله
protect(M); M.protection.autoFilter = False

# ═══ 2) القوائم (مخفي) ═══
Q = wb.create_sheet('القوائم')
for j, h in enumerate(('الغرض', 'حساب المصروف', 'حساب الأصل (20+)', 'الفئة'), 1): Q.cell(1, j, h)
for i, g in enumerate(GROUPS, 2):
    for j in range(4): Q.cell(i, j + 1, g[j])
Q.sheet_state = 'hidden'; protect(Q)

# ═══ 3) الملخص ═══
Z = wb.create_sheet('الملخص', 1); Z.sheet_view.rightToLeft = True; Z.sheet_properties.tabColor = GRN; Z.sheet_view.showGridLines = False
title(Z, 'ملخص التكاليف التشغيلية — شهر بشهر', 'A1:F1', GRN)
cell(Z, 3, 1, 'السنة', fill='C8E6C9', bold=True, size=12); cell(Z, 3, 2, 2026, '0', YEL, True, size=12)
heads = [('الشهر', 24, GRN), ('مصروف', 15, NAVY), ('أصول (%d+)' % ASSET, 15, '2F5597'), ('زكاة وصدقة', 15, 'BF8F00'), ('أجور عمال', 15, '5B3A6B'), ('المجموع', 16, '1B5E20')]
for i, (h, w, c) in enumerate(heads, 1): header(Z, 5, i, h, c, w, size=12)
Z.row_dimensions[5].height = 30
T = "'التكاليف'!"
NUM = '#,##0.00;-#,##0.00;"-"'
S = lambda m, extra: 'SUMIFS({0}{1},{0}{2},{3},{0}{4},$B$3{5})'.format(T, R('H'), R('N'), m, R('O'), extra)
for m in range(1, 13):
    r = 5 + m; band = 'FFFFFF' if m % 2 else 'F1F8F1'
    cell(Z, r, 1, '%02d — %s' % (m, MONTHS[m - 1]), fill=band, bold=True, right=True)
    cell(Z, r, 2, '=' + S(m, ',{0}{1},"مصروف",{0}{2},"مادة"'.format(T, R('J'), R('Q'))), NUM, band)             # مواد مصروف
    cell(Z, r, 3, '=' + S(m, ',{0}{1},"أصل"'.format(T, R('J'))), NUM, band)                                    # أصول (مواد + أجور ورشات)
    cell(Z, r, 4, '=' + S(m, ',{0}{1},"زكاة وصدقة"'.format(T, R('J'))), NUM, band)
    cell(Z, r, 5, '=' + S(m, ',{0}{1},"مصروف",{0}{2},"أجرة"'.format(T, R('J'), R('Q'))), NUM, band)          # أجور عمال (مش أصل)
    cell(Z, r, 6, '=SUM(B{0}:E{0})'.format(r), NUM, 'E8F5E9', True)
    Z.row_dimensions[r].height = 20
cell(Z, 18, 1, 'مجموع السنة', fill='1B5E20', bold=True, color='FFFFFF', size=12)
for c in range(2, 7): cell(Z, 18, c, '=SUM(%s6:%s17)' % ((chr(64 + c),) * 2), '#,##0.00', 'C8E6C9', True, size=12)
Z.row_dimensions[18].height = 24
Z['A19'] = 'الشهر حسب التاريخ · أجور الورشات اللي صارت أصل بعمود «أصول» · الزكاة والصدقة لحالها.'
Z['A19'].font = Font(name=FONT, size=9, italic=True, color='666666'); Z.merge_cells('A19:F19')
# حسب الغرض (بلا الورشات)
header(Z, 21, 1, 'الغرض (بلا الورشات)', GRN, size=12); header(Z, 21, 2, 'المبلغ (السنة)', GRN, size=12); Z.row_dimensions[21].height = 26
for i, g in enumerate(GROUPS):
    r = 22 + i; band = 'FFFFFF' if i % 2 else 'F1F8F1'
    cell(Z, r, 1, g[0], fill=band, bold=True, right=True)
    if g[0] == 'أخرى':   # أخرى + أي غرض مكتوب بالإيد
        f = '=SUMIFS({0}{1},{0}{2},$B$3,{0}{3},"")-SUM(B22:B{4})'.format(T, R('H'), R('O'), R('D'), r - 1)
    else:
        f = '=SUMIFS({0}{1},{0}{2},A{3},{0}{4},$B$3,{0}{5},"")'.format(T, R('H'), R('I'), r, R('O'), R('D'))
    cell(Z, r, 2, f, NUM, band)
r = 22 + NG; cell(Z, r, 1, 'المجموع', fill='1B5E20', bold=True, color='FFFFFF'); cell(Z, r, 2, '=SUM(B22:B%d)' % (r - 1), '#,##0.00', 'C8E6C9', True)
# الورشات (شغل داخلي) — بتتعبّى لحالها
WR = r + 3
x = Z.cell(WR - 1, 1, 'الورشات (شغل داخلي) — بتتعبّى لحالها'); x.font = Font(name=FONT, b=True, size=12, color=BRN)
for i, h in enumerate(('الورشة', 'مواد', 'أجور عمال', 'المجموع', 'أصل؟'), 1): header(Z, WR, i, h, BRN, size=12)
Z.row_dimensions[WR].height = 26
for k in range(1, 21):
    r = WR + k; band = 'FFFFFF' if k % 2 else 'FBF4EC'
    cell(Z, r, 1, '=IFERROR(INDEX({0}{1},MATCH({2},{0}{3},0)),"")'.format(T, R('D'), k, R('P')), fill=band, bold=True, right=True)
    cell(Z, r, 2, '=IF(A{0}="","",SUMIFS({1}{2},{1}{3},A{0},{1}{4},"مادة"))'.format(r, T, R('H'), R('D'), R('Q')), NUM, band)
    cell(Z, r, 3, '=IF(A{0}="","",SUMIFS({1}{2},{1}{3},A{0},{1}{4},"أجرة"))'.format(r, T, R('H'), R('D'), R('Q')), NUM, band)
    cell(Z, r, 4, '=IF(A{0}="","",B{0}+C{0})'.format(r), NUM, 'FBF4EC', True)
    cell(Z, r, 5, '=IF(A{0}="","",IF(D{0}>={1},"أصل","مصروف"))'.format(r, ASSET), fill=band, bold=True)
Z.conditional_formatting.add('E%d:E%d' % (WR + 1, WR + 20), FormulaRule(formula=['E%d="أصل"' % (WR + 1)], fill=F('DDEBF7'), font=Font(b=True, color='1F4E79')))
protect(Z, full=True)

# ═══ 4) التعليمات ═══
I = wb.create_sheet('التعليمات'); I.sheet_view.rightToLeft = True; I.sheet_view.showGridLines = False
I.column_dimensions['A'].width = 6; I.column_dimensions['B'].width = 22; I.column_dimensions['C'].width = 80
title(I, 'طريقة العمل', 'A1:C1', NAVY, 15, 30)
rows = [('', 'ورقة «التكاليف» — كل مصروف أو أجرة بسطر (الأعمدة: # · التاريخ · اليوم · الورشة · البند · الكمية · الإفرادي · الإجمالي · الغرض · أصل؟ · المدقق · ملاحظات)'),
        ('1', 'كل سطر: البند أو المهمة · الكمية · الإفرادي · التاريخ · الغرض · وبآخر خانة «ملاحظات إضافية»: شو هو، ليش، ولوين راح.'),
        ('2', 'الإفرادي = الكلفة (من المحل) أو سعر الشراء (من برّا).'),
        ('3', 'أجرة عامل: سطر زي أي بند — اكتب بالبند اسم العامل والشغلة (مثلاً «احمد النجار — خزائن خشب») · الكمية 1 والإفرادي = المبلغ · الغرض «أجور عمال» · بأي يوم حتى الجمعة.'),
        ('4', 'الغرض من القائمة — وإذا مش موجود اكتبه بإيدك (مثلاً «أكياس للمحل») وبينحسب «أخرى» · الفاضي بيصير أحمر و«⚠ عبّي الغرض».'),
        ('5', 'فاتورة/إيجار/زكاة: الكمية 1 والإفرادي = المبلغ · الإيجار: التحويش اليومي ما بينكتب — يوم الدفع سطر بالمبلغ كامل.'),
        ('6', 'الكمية رقم صحيح بلا فاصلة · الإفرادي والإجمالي منزلتين · الدفع نقدي — إذا تحويل كليك أو شيك اكتبها بالملاحظات.'),
        ('', 'الورشة (شغل داخلي) — لما نبني إشي للمحل (مثلاً ستاندات)'),
        ('7', 'اكتب اسم الورشة بأول عمود على كل مادة وكل أجرة إلها (والأجرة غرضها «أجور عمال») — نفس الاسم بالزبط كل مرة. فيها «ستاندات» ← الغرض «أثاث وستاندات» لحاله.'),
        ('8', 'مجموع الورشة (مواد + أجور) %d وفوق = كل أسطرها «أصل» · أقل = مصروف · الورشة بس لبناء أو شراء إشي جديد: التصليح والتعزيل والتنظيف بلا ورشة (اسم ورشة فيه «صيانة» أو «تعزيل» بيصير أحمر).' % ASSET),
        ('9', 'الفلتر على العناوين: اختار ورشة أو غرض (مثلاً «أجور عمال») وبيطلعوا أسطره لحاله.'),
        ('', 'التدقيق وأودو والملخص'),
        ('10', 'الخانة الحمرا = ناقصة · المدقق بيختار اسمه بعد ما يتأكد ← السطر بيصير أخضر · بس الأسطر المدقّقة بتنزل على أودو.'),
        ('11', 'آخر الشهر (بأمر صاحب المحل): كل سطر بينزل قيد لحاله على حسابه · المرحّل من قبل ما بيتكرّر.'),
        ('12', '«الملخص» شهر بشهر: مصروف · أصول · زكاة وصدقة · أجور عمال — وتحته حسب الغرض والورشات. الملخص والتعليمات مقفولين بكلمة السر.')]
r = 3
for n, t in rows:
    if not n:
        x = I.cell(r, 1, t); I.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        x.font = Font(name=FONT, b=True, size=12, color='FFFFFF'); x.fill = F(NAVY); x.alignment = Alignment(horizontal='right', vertical='center')
        I.row_dimensions[r].height = 24
    else:
        cell(I, r, 1, n, fill='DDE5EC', bold=True); cell(I, r, 2, t, right=True, wrap=True)
        I.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3); I.row_dimensions[r].height = 34
    r += 1
r += 1
for j, h in enumerate(('#', 'الغرض', 'أمثلة'), 1): header(I, r, j, h, GRN, size=12)
for i, g in enumerate(GROUPS, 1):
    band = 'FFFFFF' if i % 2 else 'F1F8F1'
    cell(I, r + i, 1, i, fill=band); cell(I, r + i, 2, g[0], fill=band, bold=True, right=True); cell(I, r + i, 3, g[4], fill=band, right=True, wrap=True)
    I.row_dimensions[r + i].height = 22
protect(I, full=True)
wb.calculation = CalcProperties(fullCalcOnLoad=True)   # إكسل كمبيوتر المحل بيحسب كل المعادلات أول ما ينفتح
wb.save(OUT); print('✓', OUT, wb.sheetnames)
