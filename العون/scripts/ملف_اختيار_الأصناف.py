# -*- coding: utf-8 -*-
"""ملف اختيار الأصناف — لكل بند مش مرحّل بورشة: أقرب أصناف الكتالوج، وصاحب المحل بيختار (صاحب المحل 06/10)

  python3 scripts/ملف_اختيار_الأصناف.py "ورشة_ثروت_المصري_معبّأ.xlsx" 44 "ثروت المصري"
      الملف المحلي للورشة · رقم الشريك على أودو · اسم مسودة drive/مسودات/<الاسم>.json

بيقرأ البنود اللي مش مرحّلة (حسب المسودة المحلية) · بيقترح لكل بند لحد 5 أصناف (اللي أخذه الزبون قبل أولاً)
· بيطلع `اختيار_أصناف_<الاسم>.xlsx`: عمود «اختيارك» قائمة منسدلة (رقم الخيار · جديد · يدوي) + سعر وملاحظة لصاحب المحل.
"""
import sys, os, re, json, datetime, collections
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
sys.path.insert(0, os.path.join(ROOT, 'odoo_templates'))
from jrpc import x
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation


def وحّد(s):
    s = str(s or '')
    s = re.sub('[ً-ْـ]', '', s)
    for a, b in (('أ', 'ا'), ('إ', 'ا'), ('آ', 'ا'), ('ة', 'ه'), ('ى', 'ي'), ('×', '*'), ('"', ' '), ('\\', ' ')):
        s = s.replace(a, b)
    s = re.sub(r'(\d)\s*\*\s*(\d)', r'\1*\2', s)
    return re.sub(r'\s+', ' ', s).strip().lower()

STOP = {'مع', 'من', 'الى', 'حبه', 'حجم', 'لون', 'نوع', '1', 'جودة', 'جوده', 'عاليه'}
SYN = {'ادابتر': 'ادبتر', 'سكينه': 'سكين', 'معجونه': 'معجون', 'دريل': 'درل', 'مواسير': 'ماسوره', 'نضاره': 'نظاره',
       'ص.م': 'صغير', 'شانل': 'شانيل', 'فرشاه': 'فرشايه', 'فرشاة': 'فرشايه', 'اسطوانه': 'اسطوانه'}

def كلمات(s):
    out = []
    for w in وحّد(s).replace('،', ' ').replace(',', ' ').split():
        w = SYN.get(w, w)
        if w.startswith('ال') and len(w) > 4: w = w[2:]
        if w not in STOP and (len(w) > 1 or w.isdigit()): out.append(w)
    return out

def درجة(بند, صنف):
    a, b = كلمات(بند), set(كلمات(صنف))
    if not a: return 0
    s = 0.0
    for w in a:
        if w in b: s += 2.5 if re.search(r'\d', w) else 1.0
        elif any(len(w) > 3 and (w in v or v in w) for v in b): s += 0.6
    na = {w for w in a if re.search(r'\d', w)}; nb = {w for w in b if re.search(r'\d', w)}
    if na and nb and not (na & nb): s -= 1.5          # مقاس مختلف
    return s / len(a)


def نوع_البند(بند, ن):
    b = وحّد(بند)
    if ن == 'ر': return 'مرتجع'
    if 'دفعه' in b: return 'دفعة'
    if 'كاش نقدي' in b or 'صرف نقدي' in b: return 'صرف نقدي'
    if 'فاتوره' in b: return 'فاتورة يدوية'
    return 'بند'


def اقرأ(path, مرحّل):
    """xlsx (نموذج الورشة) أو tsv (# · البند · الكمية · الإفرادي · التاريخ m/d/yyyy · المدقق · نوع ب/ر)."""
    if path.endswith('.tsv'):
        out = []
        for l in open(os.path.join(ROOT, path), encoding='utf-8').read().splitlines()[1:]:
            c = (l.split('\t') + [''] * 7)[:7]
            if not c[0]: continue
            m, d, y = (c[4].split('/') + ['', '', ''])[:3]
            try: t = datetime.date(int(y), int(m), int(d))
            except ValueError: t = c[4]
            out.append([c[0], c[1].strip(), c[2], float(c[3]) if c[3] else None, t, c[5].strip(), c[6] or 'ب'])
        return out
    ws = load_workbook(os.path.join(ROOT, path), data_only=True).worksheets[0]
    return [[r[0], str(r[1]).strip(), r[2], r[3], r[4], r[6], 'ب'] for r in ws.iter_rows(values_only=True)
            if isinstance(r[0], int) and r[0] not in مرحّل and r[0] > 43]


def main(path, pid, اسم, فقط=None, سابق=None, خيارات=5):
    مس = json.load(open(os.path.join(ROOT, 'drive', 'مسودات', اسم + '.json')))
    مرحّل = {l['#'] for l in مس['أسطر'] if l.get('حالة') == 'مرحّل'}
    قرار = {وحّد(l['بند']): l for l in مس['أسطر'] if l.get('حالة') != 'مرحّل' and l.get('كود_مقترح')}
    rows = اقرأ(path, مرحّل)
    if فقط: rows = [r for r in rows if str(r[0]) in فقط]
    سابق = json.load(open(os.path.join(ROOT, سابق))) if سابق else {}
    cat = x('product.product', 'search_read', [('sale_ok', '=', True)], ['id', 'default_code', 'name', 'list_price'])
    H = x('account.move.line', 'search_read', [('partner_id', '=', pid), ('parent_state', '=', 'posted'),
          ('display_type', '=', 'product'), ('move_id.move_type', '=', 'out_invoice')], ['product_id', 'price_unit', 'date'], order='date asc')
    hist, last = collections.Counter(), {}
    for l in H: hist[l['product_id'][0]] += 1; last[l['product_id'][0]] = l['price_unit']
    F = lambda **k: Font(name='Arial', **k); fill = lambda c: PatternFill('solid', fgColor=c)
    thin = Side(style='thin', color='D0D7E2'); B = Border(top=thin, bottom=thin, left=thin, right=thin)
    wb = Workbook(); sh = wb.active; sh.title = 'تسعير واختيار'; sh.sheet_view.rightToLeft = True
    sh.merge_cells('A1:Q1'); sh['A1'] = 'ملف تسعير واختيار الأصناف — %s · البنود المش مرحّلة (%d)' % (اسم, len(rows))
    sh['A1'].font = F(size=15, bold=True, color='FFFFFF'); sh['A1'].fill = fill('1F3A5F'); sh['A1'].alignment = Alignment(horizontal='center')
    sh.merge_cells('A2:Q2')
    sh['A2'] = ('اختار من «اختيارك» (الأصفر): 1–5 = رقم الخيار · «جديد» = صنف جديد (اسمه وسعره بالملاحظة) · «يدوي» = فاتورة يدوية سابقة · '
                '«دفعة» / «صرف» / «مرتجع» حسب النوع · عدّل «السعر المعتمد» إذا بدك. ⭐ = ثروت أخذه قبل · «ث» = آخر سعر بيع لثروت · «ك» = الكتالوج.')
    sh['A2'].font = F(size=10, color='555555'); sh['A2'].alignment = Alignment(wrap_text=True); sh.row_dimensions[2].height = 34
    hdr = ['اختيارك', '#', 'التاريخ', 'النوع', 'البند بالورقة', 'الكمية', 'سعر الورقة'] + ['خيار %d' % k for k in range(1, خيارات + 1)] + \
          ['السعر المعتمد', 'ملاحظتك', 'اختيارك السابق', 'تنبيه']
    for i, h in enumerate(hdr, 1):
        c = sh.cell(4, i, h); c.font = F(bold=True, size=10); c.fill = fill('DCE4EE'); c.border = B
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    from openpyxl.utils import get_column_letter as L_
    for k, w in enumerate([10, 6, 11, 10, 32, 8, 9] + [32] * خيارات + [10, 28, 22, 26], 1):
        sh.column_dimensions[L_(k)].width = w
    dv = DataValidation(type='list', formula1='"%s,جديد,يدوي,دفعة,صرف,مرتجع"' % ','.join(str(k) for k in range(1, خيارات + 1)), allow_blank=True); sh.add_data_validation(dv)
    لون = {'مرتجع': 'FDE2E1', 'دفعة': 'E2F0D9', 'صرف نقدي': 'E2F0D9', 'فاتورة يدوية': 'EDE7F6'}
    for i, r in enumerate(rows):
        row = 5 + i; n, بند, كم, سعر, ت, مدقق, ن = r
        نوع = نوع_البند(بند, ن)
        prev = قرار.get(وحّد(بند))
        if نوع in ('دفعة', 'صرف نقدي'): scored = []
        elif نوع == 'فاتورة يدوية': scored = [p for p in cat if p['default_code'] == '900900']
        else:
            sc = lambda p: (lambda d: d + (0.35 if hist.get(p['id']) and d > 0.3 else 0))(درجة(بند, p['name']))
            pool = [p for p in cat if hist.get(p['id'])] if نوع == 'مرتجع' else cat      # المرتجع من اللي انباع لثروت أول
            scored = [p for p in sorted(pool, key=lambda p: -sc(p))[:8] if sc(p) > 0.3]
            if نوع == 'مرتجع' and not scored:
                scored = [p for p in sorted(cat, key=lambda p: -sc(p))[:8] if sc(p) > 0.3]
            if prev:
                pp = [p for p in cat if p['default_code'] == prev['كود_مقترح']]
                scored = pp + [p for p in scored if not pp or p['id'] != pp[0]['id']]
        scored = scored[:خيارات]
        q = re.match(r'[\d.]+', str(كم)); qn = float(q.group()) if q else None
        opts, خيار_سابق = [], None
        for k, p in enumerate(scored, 1):
            opts.append('%s[%s] %s · ك %.3f%s' % ('⭐ ' if hist.get(p['id']) else '', p['default_code'], p['name'], p['list_price'],
                                                   (' · ث %.3f' % last[p['id']]) if p['id'] in last else ''))
            if prev and prev.get('كود_مقترح') == p['default_code']: خيار_سابق = k
        تنبيه = []
        سعر_م = سعر
        if سعر is None and scored:
            p0 = scored[0]; سعر_م = last.get(p0['id'], p0['list_price'])
            تنبيه.append('بلا سعر بالورقة ← ' + ('آخر سعر لثروت' if p0['id'] in last else 'سعر الكتالوج'))
        if نوع == 'مرتجع': تنبيه.append('المرتجع بسعر آخر بيع لثروت')
        if isinstance(ت, str): تنبيه.append('⚠️ التاريخ مش مقروء')
        try: tot = round(qn * float(سعر_م), 3) if qn is not None and سعر_م is not None else None
        except Exception: tot = None
        ق = سابق.get(str(n), {})
        if ق.get('سعر_معتمد') is not None: سعر_م = ق['سعر_معتمد']
        vals = [None if ق else (خيار_سابق or ({'دفعة': 'دفعة', 'صرف نقدي': 'صرف', 'فاتورة يدوية': 1}.get(نوع))), n,
                ت.strftime('%d/%m/%Y') if hasattr(ت, 'strftime') else str(ت), نوع, بند, كم, سعر] + \
               opts + [''] * (خيارات - len(opts)) + \
               [سعر_م, ق.get('ملاحظة') or '', ق.get('اختيار') or (('[%s] %s' % (prev['كود_مقترح'], prev.get('وصف') or '')) if prev else ''), ' · '.join(تنبيه)]
        for j, v in enumerate(vals, 1):
            c = sh.cell(row, j, v); c.font = F(size=10); c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
            nc = 7 + خيارات
            if j in (7, nc + 1): c.number_format = '#,##0.000'
            if نوع in لون: c.fill = fill(لون[نوع])
            if j in (1, nc + 1, nc + 2): c.fill = fill('FFF3B0')
            if j == 1: c.font = F(size=11, bold=True); c.alignment = Alignment(horizontal='center', vertical='top')
            if j == nc + 3 and v: c.fill = fill('D4EDDA')
            if j == nc + 4 and v: c.font = F(size=9, color='9C0006')
        dv.add(sh.cell(row, 1))
    sh.freeze_panes = 'F5'
    out = os.path.join(ROOT, 'اختيار_أصناف_%s.xlsx' % اسم.replace(' ', '_'))
    wb.save(out)
    print('✅ %s · %d بند' % (out, len(rows)))


if __name__ == '__main__':
    a = sys.argv[1:]
    فقط = a[a.index('--فقط') + 1].split(',') if '--فقط' in a else None
    سابق = a[a.index('--سابق') + 1] if '--سابق' in a else None
    main(a[0], int(a[1]), a[2], فقط, سابق)
