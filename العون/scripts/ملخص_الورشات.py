# -*- coding: utf-8 -*-
"""«الورشات Live» — النموذج المعتمد (صاحب المحل 06/10) · عرض فقط

  python3 scripts/ملخص_الورشات.py [ملف_الإخراج.xlsx]

3 شيتات:
  1. «ورشات الزبائن»  — ورشات الزبائن وحساباتهم الشخصية · من الأعلى مبلغ للأقل
  2. «ورشات ايلا»     — الورشات اللي مقاولها ايلا للمقاولات (طاقم م.حمزة · هشام · رامي) · مخفي وعليه باسورد
  3. «الملخص»         — مجموع الزبائن ومجموع ورشات ايلا · مخفي وعليه باسورد
الإجمالي ما بيبين إلا بالملخص. «الباقي» معادلة وبتنفحص مقابل total_due بأودو.
"""
import sys, os, datetime, collections
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
sys.path.insert(0, os.path.join(ROOT, 'odoo_templates'))
from jrpc import x
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.protection import WorkbookProtection

RECEIVABLE = 55
باسورد = __import__('os').environ.get('ALAWN_XLSX_PWD', '')
طاقم_ايلا = {6, 7, 8}          # تصنيفات أودو: م.حمزة · هشام · رامي (طاقم ايلا للمقاولات)
# بنود بالدرايف مش مرحّلة لشركاء موجودين على أودو — بتتحدّث مع كل فحص
معلّق = {48: '6 بنود (#50–55 · 04 و 06/10)'}
NUM = '#,##0.000;(#,##0.000);"-"'
أيام = ['الاثنين', 'الثلاثاء', 'الأربعاء', 'الخميس', 'الجمعة', 'السبت', 'الأحد']


def الآن():
    t = datetime.datetime.utcnow() + datetime.timedelta(hours=3)          # عمّان
    h = t.hour % 12 or 12
    return '%s %s · %d:%02d %s' % (أيام[t.weekday()], t.strftime('%d/%m/%Y'), h, t.minute,
                                     'صباحاً' if t.hour < 12 else 'مساءً')


def ارصدة():
    L = x('account.move.line', 'search_read',
          [('account_id', '=', RECEIVABLE), ('parent_state', '=', 'posted'), ('partner_id', '!=', False)],
          ['partner_id', 'debit', 'credit', 'date', 'move_type'])
    D = collections.defaultdict(lambda: {'فواتير': 0.0, 'مرتجعات': 0.0, 'دفعات': 0.0, 'صرف': 0.0,
                                         'آخر_فاتورة': '', 'آخر_دفعة': '', 'مبلغ_آخر_دفعة': 0.0})
    for l in L:
        d = D[tuple(l['partner_id'])]; t = l['move_type']
        if t == 'out_invoice':
            d['فواتير'] += l['debit'] - l['credit']; d['آخر_فاتورة'] = max(d['آخر_فاتورة'], l['date'])
        elif t == 'out_refund': d['مرتجعات'] += l['credit'] - l['debit']
        else:
            d['دفعات'] += l['credit']; d['صرف'] += l['debit']
            if l['credit'] > 0 and l['date'] >= d['آخر_دفعة']:
                d['آخر_دفعة'] = l['date']; d['مبلغ_آخر_دفعة'] = l['credit']
    ps = {p['id']: p for p in x('res.partner', 'read', [k[0] for k in D], ['total_due', 'category_id'])}
    out = []
    for (pid, name), d in D.items():
        bal = round(d['فواتير'] + d['صرف'] - d['مرتجعات'] - d['دفعات'], 3)
        if abs(bal - ps[pid]['total_due']) > 0.001:
            raise SystemExit('⛔ رصيد %s ما بطابق أودو: %.3f ≠ %.3f' % (name, bal, ps[pid]['total_due']))
        out.append(dict(id=pid, اسم=name, باقي=bal, ايلا=bool(طاقم_ايلا & set(ps[pid]['category_id'])),
                        **{k: round(v, 3) if isinstance(v, float) else v for k, v in d.items()}))
    return sorted(out, key=lambda r: -r['باقي'])


F = lambda **k: Font(name='Arial', **k)
fill = lambda c: PatternFill('solid', fgColor=c)
navy = '1F3A5F'; thin = Side(style='thin', color='D0D7E2'); B = Border(top=thin, bottom=thin, left=thin, right=thin)
HDR = ['#', 'الزبون / الورشة', 'الفواتير', 'المرتجعات', 'الدفعات', 'صرف نقدي', 'الباقي عليه',
       'تاريخ آخر فاتورة', 'تاريخ آخر دفعة', 'مبلغ آخر دفعة', 'أيام من آخر دفعة', 'بالدرايف مش مرحّل']
WIDTHS = [5, 34, 12, 11, 12, 11, 14, 14, 14, 12, 12, 26]


def رأس(ws, عنوان, وقت, cols='L'):
    ws.sheet_view.rightToLeft = True; ws.sheet_view.showGridLines = False
    ws.merge_cells(f'A1:{cols}1'); ws['A1'] = عنوان
    ws['A1'].font = F(size=18, bold=True, color='FFFFFF'); ws['A1'].fill = fill(navy)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center'); ws.row_dimensions[1].height = 34
    ws.merge_cells(f'A2:{cols}2'); ws['A2'] = 'عرض فقط · المصدر: أودو · آخر تحديث: ' + وقت
    ws['A2'].font = F(size=10, color='555555'); ws['A2'].alignment = Alignment(horizontal='center')


def بطاقة(ws, a, b, عنوان, معادلة, صيغة=None):
    ws.merge_cells(f'{a}4:{b}4'); ws.merge_cells(f'{a}5:{b}5')
    ws[f'{a}4'] = عنوان; ws[f'{a}4'].font = F(size=10, color='555555'); ws[f'{a}4'].alignment = Alignment(horizontal='center')
    ws[f'{a}5'] = معادلة; ws[f'{a}5'].font = F(size=20, bold=True, color=navy); ws[f'{a}5'].alignment = Alignment(horizontal='center')
    if صيغة: ws[f'{a}5'].number_format = صيغة
    for c in range(ord(a), ord(b) + 1):
        for r in (4, 5): ws[f'{chr(c)}{r}'].fill = fill('EEF2F7')
    ws.row_dimensions[5].height = 32


def جدول(ws, rows):
    """جدول الزبائن بدون مجاميع — يرجّع (أول صف، آخر صف)."""
    for i, w in enumerate(WIDTHS): ws.column_dimensions[chr(65 + i)].width = w
    for i, t in enumerate(HDR):
        c = ws.cell(7, i + 1, t); c.font = F(bold=True, size=10); c.fill = fill('DCE4EE')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = B
    ws.row_dimensions[7].height = 30
    s = 8
    for n, r in enumerate(rows, 1):
        row = s + n - 1; d = lambda v: datetime.date.fromisoformat(v) if v else None
        vals = [n, r['اسم'], r['فواتير'], r['مرتجعات'], r['دفعات'], r['صرف'], f'=C{row}+F{row}-D{row}-E{row}',
                d(r['آخر_فاتورة']), d(r['آخر_دفعة']), r['مبلغ_آخر_دفعة'] or None,
                f'=IF(I{row}="","لا دفعات",TODAY()-I{row})', معلّق.get(r['id'], '—')]
        for i, v in enumerate(vals):
            c = ws.cell(row, i + 1, v); c.font = F(size=10, bold=(i == 6)); c.border = B
            if i in (2, 3, 4, 5, 6, 9): c.number_format = NUM
            if i in (7, 8): c.number_format = 'dd/mm/yyyy'
            if i in (0, 7, 8, 10): c.alignment = Alignment(horizontal='center')
    e = s + len(rows) - 1; rg = f'G{s}:G{e}'
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'G{s}>500'], fill=fill('F8D7DA'), font=Font(name='Arial', color='9C0006', bold=True)))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(G{s}>=100,G{s}<=500)'], fill=fill('FFE8CC'), font=Font(name='Arial', color='9C5700', bold=True)))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'G{s}<100'], fill=fill('D4EDDA'), font=Font(name='Arial', color='1E6B34', bold=True)))
    ws.cell(e + 2, 1, '🔴 فوق 500 · 🟠 100–500 · 🟢 أقل من 100 · «صرف نقدي» = مصاري طلعت على حساب الزبون (مثل كاش للسائق) · '
            'الأرقام من القيود المرحّلة بس.').font = F(size=9, italic=True, color='777777')
    ws.freeze_panes = 'C8'
    return s, e


def شيت_زبائن(ws, عنوان, rows, وقت):
    رأس(ws, عنوان, وقت)
    s, e = s_e = (8, 7 + len(rows))
    بطاقة(ws, 'A', 'D', 'عدد الحسابات', f'=COUNTA(B{s}:B{e})')
    بطاقة(ws, 'E', 'H', 'عدد اللي عليهم رصيد', f'=COUNTIF(G{s}:G{e},">0")')
    بطاقة(ws, 'I', 'L', 'عليهم أكثر من 500', f'=COUNTIF(G{s}:G{e},">500")')
    جدول(ws, rows)
    ws.protection.sheet = True; ws.protection.set_password(باسورد)
    return s_e


def ابني(R, path):
    وقت = الآن()
    wb = Workbook()
    ws1 = wb.active; ws1.title = 'ورشات الزبائن'
    ws2 = wb.create_sheet('ورشات ايلا')
    ws3 = wb.create_sheet('الملخص')
    z = [r for r in R if not r['ايلا']]; e_ = [r for r in R if r['ايلا']]
    (a1, b1) = شيت_زبائن(ws1, 'محلات العون — ورشات الزبائن وحساباتهم الشخصية', z, وقت)
    (a2, b2) = شيت_زبائن(ws2, 'محلات العون — ورشات ايلا للمقاولات', e_, وقت)

    # الملخص — الإجماليات بس هون
    رأس(ws3, 'محلات العون — الملخص', وقت, cols='H')
    for i, w in enumerate([30, 14, 12, 12, 12, 12, 14, 12]): ws3.column_dimensions[chr(65 + i)].width = w
    h = ['', 'عدد الحسابات', 'الفواتير', 'المرتجعات', 'الدفعات', 'صرف نقدي', 'الباقي', 'النسبة']
    for i, t in enumerate(h):
        c = ws3.cell(4, i + 1, t); c.font = F(bold=True, size=11); c.fill = fill('DCE4EE'); c.border = B
        c.alignment = Alignment(horizontal='center')
    for r, (اسم, sh, a, b) in enumerate([('ورشات الزبائن وحساباتهم الشخصية', 'ورشات الزبائن', a1, b1),
                                          ('ورشات ايلا للمقاولات', 'ورشات ايلا', a2, b2)], 5):
        q = "'%s'!" % sh
        vals = [اسم, f'=COUNTA({q}B{a}:B{b})'] + [f'=SUM({q}{col}{a}:{col}{b})' for col in 'CDEFG'] + [f'=IF($G$7=0,0,G{r}/$G$7)']
        for i, v in enumerate(vals):
            c = ws3.cell(r, i + 1, v); c.font = F(size=12, bold=(i in (0, 6))); c.border = B
            if 2 <= i <= 6: c.number_format = NUM
            if i == 7: c.number_format = '0.0%'
            if i in (1, 7): c.alignment = Alignment(horizontal='center')
    vals = ['الإجمالي', '=B5+B6'] + [f'={c}5+{c}6' for c in 'CDEFG'] + ['=IF($G$7=0,0,1)']
    for i, v in enumerate(vals):
        c = ws3.cell(7, i + 1, v); c.font = F(size=13, bold=True, color='FFFFFF'); c.fill = fill(navy); c.border = B
        if 2 <= i <= 6: c.number_format = NUM
        if i == 7: c.number_format = '0.0%'
        if i in (1, 7): c.alignment = Alignment(horizontal='center')
    ws3.row_dimensions[7].height = 28
    ws3.cell(9, 1, 'الباقي = الفواتير + صرف نقدي − المرتجعات − الدفعات · مطابق لرصيد أودو لكل حساب.').font = F(size=9, italic=True, color='777777')
    ws3.protection.sheet = True; ws3.protection.set_password(باسورد)

    # الشيتين المحميات: مخفيين + قفل ترتيب الملف بالباسورد (ما بينفتحوا إلا بالباسورد)
    ws2.sheet_state = 'hidden'; ws3.sheet_state = 'hidden'
    wb.security = WorkbookProtection(workbookPassword=باسورد, lockStructure=True)
    wb.active = 0
    wb.save(path)
    return len(z), len(e_)


if __name__ == '__main__':
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'الورشات_Live.xlsx')
    R = ارصدة(); nz, ne = ابني(R, out)
    print('✅ %s' % out)
    print('الزبائن: %d · %.3f' % (nz, sum(r['باقي'] for r in R if not r['ايلا'])))
    print('ايلا: %d · %.3f  ←  %s' % (ne, sum(r['باقي'] for r in R if r['ايلا']), ' · '.join(r['اسم'] for r in R if r['ايلا'])))
    print('الإجمالي: %.3f' % sum(r['باقي'] for r in R))
