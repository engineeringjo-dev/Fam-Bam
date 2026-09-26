# -*- coding: utf-8 -*-
"""نموذج بيانات الموظفين — محلات العون لمواد البناء
المتطلبات الثلاث لتشغيل «مدير الموظفين»:
  ١) الموظفون وأنواعهم وأجورهم ويوم القبض   ٢) الأرصدة الافتتاحية (سلف · ديون · بضاعة)   ٣) السقوف
+ ورقة «المهام» (أسعار المهام المتكررة لعمال حسب الطلب) + التعليمات.
الخلايا الصفراء هي اللي بتتعبّى. الرمادية بتنحسب لحالها.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

OUT = '/home/user/Fam-Bam/العون/نموذج_الموظفين.xlsx'
NAVY, TEAL, GOLD, PLUM = '1F4E6B', '0F6E6E', '8A6A00', '5B3A6B'
EDIT, AUTO, GREY = 'FFF6CC', 'EDEDED', 'F5F5F5'
LINE = Side(style='thin', color='9AA5AD')
THICK = Side(style='medium', color=NAVY)
BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
R0 = 6                      # أول سطر بيانات
EMP_ROWS, BAL_ROWS, TASK_ROWS = 20, 40, 30
EMP_END, BAL_END = R0 + EMP_ROWS - 1, R0 + BAL_ROWS - 1

def F(sz=11, b=False, c='000000'):
    return Font(name='Arial', size=sz, bold=b, color=c)
def C(h='center'):
    return Alignment(horizontal=h, vertical='center', wrap_text=True)
def fill(c):
    return PatternFill('solid', fgColor=c)

wb = openpyxl.Workbook()

def sheet(title, banner, color, headers, widths, nrows, note=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    n = len(headers)
    last = openpyxl.utils.get_column_letter(n)
    ws.merge_cells('A1:%s1' % last)
    ws['A1'] = banner
    ws['A1'].font = F(15, True, 'FFFFFF'); ws['A1'].alignment = C(); ws['A1'].fill = fill(color)
    ws.row_dimensions[1].height = 30
    ws.merge_cells('A2:%s2' % last)
    ws['A2'] = note or ''
    ws['A2'].font = F(10, False, '444444'); ws['A2'].alignment = C('right')
    ws.row_dimensions[2].height = 34
    ws.merge_cells('A3:%s3' % last)
    ws['A3'] = '🟨 خلية صفراء = تتعبّى   ⬜ رمادية = تنحسب لحالها   ⚪ بيضاء = اختياري'
    ws['A3'].font = F(10, True, GOLD); ws['A3'].alignment = C('right')
    for k, w in zip(range(1, n + 1), widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(k)].width = w
    for i, h in enumerate(headers, start=1):
        c = ws.cell(5, i, h)
        c.font = F(11, True, 'FFFFFF'); c.alignment = C(); c.fill = fill(color)
        c.border = Border(left=LINE, right=LINE, top=THICK, bottom=THICK)
    ws.row_dimensions[5].height = 30
    for r in range(R0, R0 + nrows):
        ws.row_dimensions[r].height = 22
        for col in range(1, n + 1):
            c = ws.cell(r, col)
            c.border = BOX; c.font = F(11); c.alignment = C(); c.fill = fill('FFFFFF')
    ws.freeze_panes = 'A6'
    return ws

def paint(ws, cols, lo, hi, color):
    for col in cols:
        for r in range(lo, hi + 1):
            ws['%s%d' % (col, r)].fill = fill(color)

def dv(ws, formula, rng, prompt=None):
    d = DataValidation(type='list', formula1=formula, allow_blank=True)
    if prompt:
        d.promptTitle = 'اختر'; d.prompt = prompt; d.showInputMessage = True
    ws.add_data_validation(d); d.add(rng)

# ═══════════ ١) الموظفون ═══════════
E = sheet('الموظفين', '١) الموظفون — محلات العون لمواد البناء', NAVY,
          ['#', 'الاسم', 'النوع', 'الأجر (دينار)', 'الأجر لكل', 'يوم القبض', 'الهاتف', 'تاريخ البدء', 'ملاحظات'],
          [5, 24, 16, 13, 12, 14, 15, 13, 34], EMP_ROWS,
          'سطر لكل موظف. عمال المهام: سطر لكل عامل بالاسم (الأجر يُترك فاضياً — كل مهمة بسعرها من ورقة «المهام»).')
paint(E, 'BCDEFGH', R0, EMP_END, EDIT)
paint(E, 'A', R0, EMP_END, AUTO)
for r in range(R0, R0 + EMP_ROWS):
    E['A%d' % r] = '=IF(B%d="","",ROW()-%d)' % (r, R0 - 1)
    E['D%d' % r].number_format = '#,##0.000'
    E['H%d' % r].number_format = 'DD/MM/YYYY'
dv(E, '"دائم أسبوعي,دائم كل 10 أيام,مياومة,عامل مهام"', 'C%d:C%d' % (R0, EMP_END), 'دائم أسبوعي · دائم كل 10 أيام · مياومة · عامل مهام')
dv(E, '"أسبوع,10 أيام,يوم,مهمة"', 'E%d:E%d' % (R0, EMP_END))
dv(E, '"السبت,الأحد,الاثنين,الثلاثاء,الأربعاء,الخميس,الجمعة,آخر يوم دوام,عند التسليم"', 'F%d:F%d' % (R0, EMP_END))
# المعروف مسبقاً من كلام صاحب المحل
pre = [('المدير', 'دائم كل 10 أيام', None, '10 أيام', None, 'الأجر ويوم القبض'),
       ('عمر المصري', 'دائم أسبوعي', None, 'أسبوع', None, 'الأجر ويوم القبض'),
       ('عبدالعزيز', 'مياومة', None, 'يوم', 'آخر يوم دوام', 'أجر اليوم — يقبض نهاية كل يوم بيجي فيه'),
       ('', 'عامل مهام', None, 'مهمة', 'عند التسليم', 'اكتب اسم كل عامل مهام بسطر'),
       ('', 'عامل مهام', None, 'مهمة', 'عند التسليم', '')]
for i, (n, t, w, per, day, note) in enumerate(pre):
    r = R0 + i
    E['B%d' % r] = n; E['C%d' % r] = t; E['E%d' % r] = per
    if day: E['F%d' % r] = day
    E['I%d' % r] = note
    E['I%d' % r].font = F(10, False, '666666'); E['I%d' % r].alignment = C('right')
E['D%d' % R0].comment = Comment('المطلوب: أجر الـ10 أيام كاملة (مش شهري ومش يومي).', 'مدير الموظفين')
E['D%d' % (R0 + 1)].comment = Comment('المطلوب: أجر الأسبوع كاملاً.', 'مدير الموظفين')
E['D%d' % (R0 + 2)].comment = Comment('المطلوب: أجر اليوم الواحد.', 'مدير الموظفين')
# مثال توضيحي أسفل الجدول
ex = R0 + EMP_ROWS + 1
E.merge_cells('B%d:I%d' % (ex, ex))
E['B%d' % ex] = 'مثال للتعبئة (مش بيانات حقيقية): «عمر المصري · دائم أسبوعي · 100.000 · أسبوع · الخميس · 079xxxxxxx · 01/01/2026»'
E['B%d' % ex].font = F(10, False, '666666'); E['B%d' % ex].alignment = C('right')

# ═══════════ ٢) الأرصدة الافتتاحية ═══════════
B = sheet('الأرصدة الافتتاحية', '٢) الأرصدة الافتتاحية — شو على كل موظف اليوم', TEAL,
          ['#', 'الموظف', 'النوع', 'التاريخ', 'الأصل (دينار)', 'المسدَّد لحد اليوم', 'المتبقي', 'القسط', 'القسط كل', 'ملاحظات / وصف البضاعة'],
          [5, 22, 12, 13, 14, 15, 13, 12, 12, 40], BAL_ROWS,
          'سطر لكل سلفة أو دين أو بضاعة قائمة الآن. السلفة العادية اللي بتنخصم كلها آخر الأسبوع: القسط = الأصل. البضاعة: اكتب وصفها بالملاحظات وسعرها اللي انحسب عليه.')
paint(B, 'BCDEFHIJ', R0, BAL_END, EDIT)
paint(B, 'AG', R0, BAL_END, AUTO)
for r in range(R0, R0 + BAL_ROWS):
    B['A%d' % r] = '=IF(B%d="","",ROW()-%d)' % (r, R0 - 1)
    B['G%d' % r] = '=IF(E%d="","",E%d-F%d)' % (r, r, r)
    for col in 'EFGH':
        B['%s%d' % (col, r)].number_format = '#,##0.000'
    B['D%d' % r].number_format = 'DD/MM/YYYY'
dv(B, "='الموظفين'!$B$%d:$B$%d" % (R0, EMP_END), 'B%d:B%d' % (R0, BAL_END), 'الأسماء من ورقة الموظفين')
dv(B, '"سلفة,دين,بضاعة"', 'C%d:C%d' % (R0, BAL_END))
dv(B, '"أسبوع,10 أيام,يوم,دفعة وحدة"', 'I%d:I%d' % (R0, BAL_END))
# تحذير لو القسط أكبر من المتبقي
B.conditional_formatting.add('H%d:H%d' % (R0, BAL_END),
    FormulaRule(formula=['AND(H%d<>"",G%d<>"",H%d>G%d)' % (R0, R0, R0, R0)], fill=fill('FBD5D5')))
# ملخص لكل موظف
S0 = R0 + BAL_ROWS + 2
B.merge_cells('B%d:F%d' % (S0, S0))
B['B%d' % S0] = 'ملخص المتبقي على كل موظف (ينحسب لحاله)'
B['B%d' % S0].font = F(11, True, 'FFFFFF'); B['B%d' % S0].fill = fill(TEAL); B['B%d' % S0].alignment = C()
for i, h in enumerate(['الموظف', 'سلف', 'ديون', 'بضاعة', 'المجموع'], start=2):
    c = B.cell(S0 + 1, i, h); c.font = F(10, True); c.fill = fill(GREY); c.border = BOX; c.alignment = C()
for k in range(EMP_ROWS):
    r = S0 + 2 + k; er = R0 + k
    B['B%d' % r] = "=IF('الموظفين'!B%d=\"\",\"\",'الموظفين'!B%d)" % (er, er)
    for col, typ in zip('CDE', ['سلفة', 'دين', 'بضاعة']):
        B['%s%d' % (col, r)] = '=IF($B%d="","",SUMIFS($G$%d:$G$%d,$B$%d:$B$%d,$B%d,$C$%d:$C$%d,"%s"))' % (
            r, R0, BAL_END, R0, BAL_END, r, R0, BAL_END, typ)
    B['F%d' % r] = '=IF(B%d="","",SUM(C%d:E%d))' % (r, r, r)
    for col in 'BCDEF':
        c = B['%s%d' % (col, r)]; c.border = BOX; c.font = F(10); c.alignment = C(); c.fill = fill(AUTO)
        if col != 'B': c.number_format = '#,##0.000;[Red]-#,##0.000;-'
ex = S0 + 2 + EMP_ROWS + 1
B.merge_cells('B%d:J%d' % (ex, ex))
B['B%d' % ex] = 'مثال للتعبئة: «عمر المصري · دين · 01/09/2026 · 200.000 · 25.000 · (يُحسب 175.000) · 25.000 · أسبوع · دين كاش»'
B['B%d' % ex].font = F(10, False, '666666'); B['B%d' % ex].alignment = C('right')

# ═══════════ ٣) السقوف ═══════════
K = wb.create_sheet('السقوف')
K.sheet_view.rightToLeft = True; K.sheet_view.showGridLines = False
K.merge_cells('A1:E1'); K['A1'] = '٣) السقوف والقواعد'
K['A1'].font = F(15, True, 'FFFFFF'); K['A1'].fill = fill(GOLD); K['A1'].alignment = C()
K.row_dimensions[1].height = 30
K.column_dimensions['A'].width = 4; K.column_dimensions['B'].width = 44
K.column_dimensions['C'].width = 16; K.column_dimensions['D'].width = 16; K.column_dimensions['E'].width = 40
for i, h in enumerate(['#', 'القاعدة', 'القيمة', 'الوحدة', 'الشرح'], start=1):
    c = K.cell(3, i, h); c.font = F(11, True, 'FFFFFF'); c.fill = fill(GOLD); c.alignment = C()
    c.border = Border(left=LINE, right=LINE, top=THICK, bottom=THICK)
rules = [
    ('سقف السلفة خلال الدورة', 50, '% من أجر الدورة', 'أقصى مجموع سلف يقدر ياخذها الموظف بين قبضتين. اكتب 0 = بلا سقف (حسب تقديرك كل مرة).'),
    ('سقف مجموع الأقساط بالدورة', 40, '% من أجر الدورة', 'أقصى نسبة تنخصم أقساطاً (دين + بضاعة) من قبضة وحدة عشان يضل يقبض شي.'),
    ('سقف الدين الإجمالي للموظف', 0, 'دينار', 'أقصى رصيد مدين مسموح (دين + بضاعة قائمة). 0 = بلا سقف.'),
    ('سعر البضاعة للموظف', 'سعر الجملة', 'اختيار', 'سعر الجملة · سعر التكلفة · سعر الكتالوج · حسب كل حالة'),
    ('السلفة بتنخصم', 'كاملة بنفس الدورة', 'اختيار', 'كاملة بنفس الدورة · أو ممكن تتقسّط لو كبيرة'),
    ('إذا الصافي طلع سالب (السلف أكثر من الأجر)', 'يترحّل للدورة الجاية', 'اختيار', 'يترحّل للدورة الجاية · يُحوَّل لدين بقسط'),
    ('يوم إقفال الدورة والكشف', 'الخميس', 'يوم', 'اليوم اللي بيطلعلك فيه مدير الموظفين جدول الصافي لكل موظف'),
]
for i, (name, val, unit, desc) in enumerate(rules, start=1):
    r = 3 + i
    K.cell(r, 1, i); K.cell(r, 2, name); K.cell(r, 3, val); K.cell(r, 4, unit); K.cell(r, 5, desc)
    for col in range(1, 6):
        c = K.cell(r, col); c.border = BOX; c.font = F(11); c.alignment = C('right' if col in (2, 5) else 'center')
    K.cell(r, 3).fill = fill(EDIT); K.cell(r, 3).font = F(11, True)
    K.cell(r, 5).font = F(10, False, '555555')
    K.row_dimensions[r].height = 34
dv(K, '"سعر الجملة,سعر التكلفة,سعر الكتالوج,حسب كل حالة"', 'C7')
dv(K, '"كاملة بنفس الدورة,ممكن تتقسّط"', 'C8')
dv(K, '"يترحّل للدورة الجاية,يُحوَّل لدين بقسط"', 'C9')
dv(K, '"السبت,الأحد,الاثنين,الثلاثاء,الأربعاء,الخميس,الجمعة"', 'C10')
K['B12'] = 'القيم المكتوبة هي اقتراحي — غيّرها براحتك. 0 = بلا سقف.'
K['B12'].font = F(10, True, GOLD); K['B12'].alignment = C('right')

# ═══════════ المهام ═══════════
T = sheet('المهام', 'أسعار المهام المتكررة — لعمال حسب الطلب', PLUM,
          ['#', 'المهمة', 'السعر (دينار)', 'الوحدة', 'ملاحظات'], [5, 40, 14, 14, 40], TASK_ROWS,
          'قائمة المهام اللي بتتكرر وسعر كل وحدة، عشان لما تقول «تنزيل طبلية جبسمبورد – فلان» ينزل السطر بسعره لحاله. المهمة الجديدة بتنضاف وقتها.')
paint(T, 'BCDE', R0, R0 + TASK_ROWS - 1, EDIT)
paint(T, 'A', R0, R0 + TASK_ROWS - 1, AUTO)
for r in range(R0, R0 + TASK_ROWS):
    T['A%d' % r] = '=IF(B%d="","",ROW()-%d)' % (r, R0 - 1)
    T['C%d' % r].number_format = '#,##0.000'
dv(T, '"طبلية,شوال,حبة,ساعة,يوم,نقلة,مقطوعية"', 'D%d:D%d' % (R0, R0 + TASK_ROWS - 1))
tasks = [('تنزيل طبلية جبسمبورد', None, 'طبلية'), ('تنزيل طبلية اسمنت', None, 'طبلية'),
         ('تحميل / تنزيل شوالات', None, 'شوال'), ('توصيل بضاعة داخل المدينة', None, 'نقلة'),
         ('تنظيف المستودع', None, 'مقطوعية')]
for i, (n, p, u) in enumerate(tasks):
    T['B%d' % (R0 + i)] = n; T['D%d' % (R0 + i)] = u
T['B%d' % (R0 + TASK_ROWS + 1)] = 'المهام المكتوبة أمثلة من كلامك — عدّل أو احذف أو أضف.'
T['B%d' % (R0 + TASK_ROWS + 1)].font = F(10, False, '666666'); T['B%d' % (R0 + TASK_ROWS + 1)].alignment = C('right')

# ═══════════ التعليمات ═══════════
I = wb.create_sheet('التعليمات')
I.sheet_view.rightToLeft = True; I.sheet_view.showGridLines = False
I.column_dimensions['A'].width = 5; I.column_dimensions['B'].width = 110
I['B1'] = 'كيف تتعبّى'; I['B1'].font = F(15, True, NAVY)
steps = [
    ('ورقة «الموظفين»', None),
    ('1', 'الأسماء الثلاثة مكتوبة من كلامك. الناقص: الأجر (الصفراء) ويوم القبض والهاتف. عمال المهام: اكتب اسم كل واحد بسطر — أجره فاضي لأنه بيتحاسب بالمهمة.'),
    ('2', 'الأجر يُكتب لكامل الدورة: أجر الـ10 أيام للمدير · أجر الأسبوع لعمر · أجر اليوم لعبدالعزيز.'),
    ('ورقة «الأرصدة الافتتاحية»', None),
    ('3', 'كل شي قائم اليوم على أي موظف، سطر لحاله: سلفة · دين · بضاعة. الأصل والمسدَّد لحد اليوم، والمتبقي بينحسب. القسط وكل قديش.'),
    ('4', 'البضاعة: اكتب وصفها وسعرها اللي انحسب عليه بالملاحظات. لو بدك، ابعتلي فاتورتها وبدخّلها كفاتورة بيع عليه.'),
    ('5', 'الملخص تحت الجدول بيجمعلك لحاله شو على كل واحد.'),
    ('ورقة «السقوف»', None),
    ('6', 'القيم اقتراحي. غيّرها أو اكتب 0 = بلا سقف. هاي القواعد اللي بيمشي عليها مدير الموظفين تلقائياً وبينبّهك لما تنكسر.'),
    ('ورقة «المهام»', None),
    ('7', 'أسعار المهام المتكررة. لما تكتبلي «تنزيل طبلية جبسمبورد – محمد» بينزل السطر بسعره من هون.'),
    ('بعدها', None),
    ('8', 'ابعتلي الملف. بفتح لكل موظف شريكاً وسجلّاً على أودو، وبدخّل الأرصدة الافتتاحية كمسودة، وبعرضها عليك قبل الترحيل.'),
]
r = 3
for a, b in steps:
    if b is None:
        I['B%d' % r] = a; I['B%d' % r].font = F(12, True, TEAL); r += 1; continue
    I['A%d' % r] = a; I['A%d' % r].font = F(11, True); I['A%d' % r].alignment = C()
    I['B%d' % r] = b; I['B%d' % r].font = F(11); I['B%d' % r].alignment = C('right')
    I.row_dimensions[r].height = 34; r += 1

del wb['Sheet']
wb.save(OUT)
print('saved', OUT)
