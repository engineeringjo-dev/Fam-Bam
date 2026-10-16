# -*- coding: utf-8 -*-
"""نموذج مشتريات شهري — محلات العون لمواد البناء
ورقة «الفواتير» (سطر لكل فاتورة) + ورقة «البنود» (للفواتير اليدوية فقط) + التعليمات.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'نموذج_مشتريات_شهري.xlsx')
INV_ROWS, ITEM_ROWS = 200, 600
R0 = 7
INV_END, ITEM_END = R0 + INV_ROWS - 1, R0 + ITEM_ROWS - 1

NAVY, TEAL, GOLD = '1F4E6B', '0F6E6E', '8A6A00'
EDIT, AUTO = 'FFF6CC', 'EDEDED'
GREEN, RED = 'D6F0D6', 'FBD5D5'
LINE  = Side(style='thin', color='9AA5AD')
THICK = Side(style='medium', color=NAVY)
BOX   = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)

def F(sz=11, b=False, c='000000'):
    return Font(name='Arial', size=sz, bold=b, color=c)
def C(h='center'):
    return Alignment(horizontal=h, vertical='center', wrap_text=True)

wb = openpyxl.Workbook()

# مجموع بنود فاتورة معيّنة من ورقة البنود
SUMIT = "SUMIF('البنود'!$B${0}:$B${1},D{{r}},'البنود'!$F${0}:$F${1})".format(R0, ITEM_END)


def head(ws, banner, color, ncols, label, value_formula=None):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    last = chr(ord('A') + ncols - 1)
    ws.merge_cells('A1:%s1' % last)
    ws['A1'] = banner
    ws['A1'].font = F(16, True, 'FFFFFF'); ws['A1'].alignment = C()
    ws['A1'].fill = PatternFill('solid', fgColor=color)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 26
    ws['A2'] = label
    ws['A2'].font = F(10, True); ws['A2'].alignment = C('right')
    ws['A2'].fill = PatternFill('solid', fgColor='E8EDF0'); ws['A2'].border = BOX
    ws.merge_cells('B2:%s2' % last)
    if value_formula:
        ws['B2'] = value_formula
        ws['B2'].fill = PatternFill('solid', fgColor=AUTO)
    else:
        ws['B2'].fill = PatternFill('solid', fgColor=EDIT)
    ws['B2'].font = F(13, True); ws['B2'].alignment = C()
    for col in range(2, ncols + 1):
        ws.cell(2, col).border = BOX


def table(ws, headers, widths, color, nrows):
    for k, v in zip([chr(ord('A') + i) for i in range(len(widths))], widths):
        ws.column_dimensions[k].width = v
    for i, h in enumerate(headers, start=1):
        c = ws.cell(6, i, h)
        c.font = F(11, True, 'FFFFFF'); c.alignment = C()
        c.fill = PatternFill('solid', fgColor=color)
        c.border = Border(left=LINE, right=LINE, top=THICK, bottom=THICK)
    ws.row_dimensions[6].height = 28
    for r in range(R0, R0 + nrows):
        ws.row_dimensions[r].height = 21
        for col in range(1, len(headers) + 1):
            c = ws.cell(r, col)
            c.border = BOX; c.font = F(11); c.alignment = C()
            c.fill = PatternFill('solid', fgColor='FFFFFF')


def lock(ws, ncols, last_row, unlocked):
    for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=ncols):
        for c in row:
            c.protection = Protection(locked=True)
    for col, lo, hi in unlocked:
        for r in range(lo, hi + 1):
            ws.cell(r, col).protection = Protection(locked=False)
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.selectLockedCells = False


def page(ws, ncols, last_row):
    last = chr(ord('A') + ncols - 1)
    ws.print_area = 'A1:%s%d' % (last, last_row)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_title_rows = '1:6'
    ws.freeze_panes = 'A7'


# ══════════ ورقة الفواتير ══════════
inv = wb.active
inv.title = 'الفواتير'
head(inv, 'مشتريات البضاعة — محلات العون لمواد البناء', NAVY, 12, 'الشهر')

inv.row_dimensions[4].height = 26
for cell, txt in (('A4', 'عدد الفواتير'), ('C4', 'مجموع الصافي')):
    inv[cell] = txt
    inv[cell].font = F(10, True); inv[cell].alignment = C('right')
    inv[cell].fill = PatternFill('solid', fgColor='E8EDF0'); inv[cell].border = BOX
inv['B4'] = '=COUNTA(B{}:B{})'.format(R0, INV_END)
inv['D4'] = '=SUM(H{}:H{})'.format(R0, INV_END)
inv['D4'].number_format = '0.000'
for c in ('B4', 'D4'):
    inv[c].fill = PatternFill('solid', fgColor=AUTO)
    inv[c].font = F(11, True); inv[c].alignment = C(); inv[c].border = BOX
inv.merge_cells('E4:L4')
inv['E4'] = ('=IF(B4=0,"⛔ ما في فواتير مُدخلة",'
             'IF(COUNTIF(I{0}:I{1},"⛔*")+COUNTIF(I{0}:I{1},"🟡*")>0,'
             '"⛔ في "&(COUNTIF(I{0}:I{1},"⛔*")+COUNTIF(I{0}:I{1},"🟡*"))&" فاتورة ناقصة",'
             '"✅ كل الفواتير مكتملة"))').format(R0, INV_END)
inv['E4'].fill = PatternFill('solid', fgColor=AUTO)
inv['E4'].font = F(12, True); inv['E4'].alignment = C(); inv['E4'].border = BOX

table(inv,
      ['#', 'المورد', 'التاريخ', 'رقم الفاتورة', 'النوع',
       'المجموع', 'الخصم', 'الصافي', 'الحالة', 'اسم الصورة', 'الصورة (Ctrl+K)', '✓'],
      [5, 30, 16, 14, 12, 12, 10, 12, 26, 26, 22, 6], NAVY, INV_ROWS)

for r in range(R0, INV_END + 1):
    inv.cell(r, 1, '=IF(B{}="","",ROW()-{})'.format(r, R0 - 1))
    inv.cell(r, 1).fill = PatternFill('solid', fgColor=AUTO)
    inv.cell(r, 2).alignment = C('right')
    inv.cell(r, 3).number_format = '00"/"00"/"0000'
    for col in (6, 7, 8):
        inv.cell(r, col).number_format = '0.000'
    inv.cell(r, 8, '=IF(F{0}="","",F{0}-N(G{0}))'.format(r))
    inv.cell(r, 8).fill = PatternFill('solid', fgColor=AUTO)
    s = SUMIT.format(r=r)
    inv.cell(r, 9, (
        '=IF(B{r}="","",'
        'IF(E{r}="","⛔ اختر النوع",'
        'IF(E{r}="مطبوعة",IF(K{r}="","⛔ ارفق صورتها","📷 مرفقة"),'
        'IF(D{r}="","⛔ اكتب رقم الفاتورة",'
        'IF({s}=0,"🟡 بنودها ما انزلت",'
        'IF(ROUND({s}-N(F{r}),3)=0,"✅ خلصت — البنود مطابقة",'
        '"⛔ فرق "&TEXT(N(F{r})-{s},"0.000")))))))'
    ).format(r=r, s=s))
    inv.cell(r, 9).fill = PatternFill('solid', fgColor=AUTO)
    inv.cell(r, 9).font = F(10, True)
    # صور الفواتير المطبوعة (13/10): الاسم جاهز «المورد_رقم الفاتورة.jpg» · الموظف بيوقف على K وبيكبس Ctrl+K وبيختار الصورة
    inv.cell(r, 10, ('=IF(AND(B{r}<>"",E{r}="مطبوعة",D{r}<>""),'
                     'SUBSTITUTE(SUBSTITUTE(TRIM(B{r})&"_"&TRIM(D{r}),"/","-"),"\\","-")&".jpg","")').format(r=r))
    inv.cell(r, 10).fill = PatternFill('solid', fgColor=AUTO)
    inv.cell(r, 10).font = F(10)
    inv.cell(r, 11).font = Font(name='Arial', size=10, color='0563C1', underline='single')
    inv.cell(r, 12, '=IF(OR(B{r}="",E{r}<>"مطبوعة"),"",IF(K{r}<>"","✓",""))'.format(r=r))
    inv.cell(r, 12).fill = PatternFill('solid', fgColor=AUTO)
    inv.cell(r, 12).font = F(14, True, '17632A')

d = DataValidation(type='list', formula1='"مطبوعة,يدوية"', allow_blank=True, showDropDown=False)
inv.add_data_validation(d); d.add('E{}:E{}'.format(R0, INV_END))

# «يدوية» ← خيارات وأعمدة الصورة بتختفي (14/10): خانات J:L بتصير رمادي غامق بلا خطوط، وخانة الصورة ما بتقبل كتابة
inv.conditional_formatting.add('J{}:L{}'.format(R0, INV_END), FormulaRule(
    formula=['$E{}="يدوية"'.format(R0)], fill=PatternFill('solid', bgColor='BFBFBF'),
    font=Font(name='Arial', size=10, color='BFBFBF'),
    border=Border(left=Side(style=None), right=Side(style=None), top=Side(style=None), bottom=Side(style=None)),
    stopIfTrue=True))
dk = DataValidation(type='custom', formula1='$E{}<>"يدوية"'.format(R0), allow_blank=True)
dk.error = 'الفاتورة اليدوية ما إلها صورة — بنودها بورقة «البنود».'
dk.errorTitle = 'يدوية'
dk.showErrorMessage = True
inv.add_data_validation(dk); dk.add('K{}:K{}'.format(R0, INV_END))

inv.conditional_formatting.add('E4', FormulaRule(
    formula=['LEFT(E4,1)="✅"'], fill=PatternFill('solid', bgColor=GREEN), font=F(12, True, '17632A')))
inv.conditional_formatting.add('E4', FormulaRule(
    formula=['LEFT(E4,1)="⛔"'], fill=PatternFill('solid', bgColor=RED), font=F(12, True, 'A11212')))
rng = 'I{}:I{}'.format(R0, INV_END)
inv.conditional_formatting.add(rng, FormulaRule(
    formula=['LEFT(I{},1)="⛔"'.format(R0)], fill=PatternFill('solid', bgColor=RED), font=F(10, True, 'A11212')))
inv.conditional_formatting.add(rng, FormulaRule(
    formula=['LEFT(I{},1)="✅"'.format(R0)], fill=PatternFill('solid', bgColor=GREEN), font=F(10, True, '17632A')))
# رقم فاتورة مكرر — بيخرّب ربط البنود
inv.conditional_formatting.add('D{}:D{}'.format(R0, INV_END), FormulaRule(
    formula=['AND($D{0}<>"",COUNTIF($D${1}:$D${2},$D{0})>1)'.format(R0, R0, INV_END)],
    fill=PatternFill('solid', bgColor='FFD9A0')))
# تاريخ غير منطقي — يوم>31 أو شهر>12 أو سنة<2020
BADDATE = ('AND($B{0}<>"",$C{0}<>"",OR(INT($C{0}/1000000)<1,INT($C{0}/1000000)>31,'
           'MOD(INT($C{0}/10000),100)<1,MOD(INT($C{0}/10000),100)>12,'
           'MOD($C{0},10000)<2020))').format(R0)
inv.conditional_formatting.add('C{}:C{}'.format(R0, INV_END),
    FormulaRule(formula=[BADDATE], fill=PatternFill('solid', bgColor='FFD9A0'),
                font=F(11, True, 'A11212')))

# حقول ناقصة
for col in ('C', 'F'):
    inv.conditional_formatting.add('{0}{1}:{0}{2}'.format(col, R0, INV_END), FormulaRule(
        formula=['AND($B{0}<>"",${1}{0}="")'.format(R0, col)],
        fill=PatternFill('solid', bgColor=RED)))

W = INV_END + 2
inv.merge_cells('A{0}:L{0}'.format(W))
inv['A{}'.format(W)] = 'المطبوعة: تاريخ + مورد + رقم + مجموع، وصورتها بمجلّد «صور الفواتير» باسم عمود «اسم الصورة» ثم Ctrl+K على خانة الصورة. اليدوية: بنودها بورقة «البنود».'
inv['A{}'.format(W)].font = F(10, True, 'FFFFFF'); inv['A{}'.format(W)].alignment = C()
inv['A{}'.format(W)].fill = PatternFill('solid', fgColor=TEAL)
inv.row_dimensions[W].height = 22

page(inv, 12, W)
lock(inv, 12, W, [(c, R0, INV_END) for c in (2, 3, 4, 5, 6, 7, 11)] + [(2, 2, 2)])
inv.protection.insertHyperlinks = False   # Ctrl+K مسموح على خانة الصورة

# ══════════ ورقة البنود ══════════
it = wb.create_sheet('البنود')
head(it, 'بنود الفواتير اليدوية — محلات العون لمواد البناء', TEAL, 6, 'الشهر',
     value_formula="=IF('الفواتير'!B2=\"\",\"\",'الفواتير'!B2)")

it.row_dimensions[4].height = 26
for cell, txt in (('A4', 'عدد البنود'), ('C4', 'مجموعها')):
    it[cell] = txt
    it[cell].font = F(10, True); it[cell].alignment = C('right')
    it[cell].fill = PatternFill('solid', fgColor='E8EDF0'); it[cell].border = BOX
it['B4'] = '=COUNTA(C{}:C{})'.format(R0, ITEM_END)
it['D4'] = '=SUM(F{}:F{})'.format(R0, ITEM_END)
it['D4'].number_format = '0.000'
for c in ('B4', 'D4'):
    it[c].fill = PatternFill('solid', fgColor=AUTO)
    it[c].font = F(11, True); it[c].alignment = C(); it[c].border = BOX
it.merge_cells('E4:F4')
it['E4'] = ('=IF(B4=0,"لا يوجد بنود",'
            'IF(COUNTIF(A{0}:A{1},"⛔*")>0,"⛔ في بند رقم فاتورته مش موجود","✅ كل البنود مربوطة"))'
            ).format(R0, ITEM_END)
it['E4'].fill = PatternFill('solid', fgColor=AUTO)
it['E4'].font = F(12, True); it['E4'].alignment = C(); it['E4'].border = BOX

table(it, ['#', 'رقم الفاتورة', 'البند', 'الكمية', 'الإفرادي', 'الإجمالي'],
      [7, 14, 50, 10, 12, 13], TEAL, ITEM_ROWS)

for r in range(R0, ITEM_END + 1):
    it.cell(r, 1, ('=IF(C{r}="","",'
                   'IF(COUNTIF(\'الفواتير\'!$D${0}:$D${1},B{r})=0,"⛔",ROW()-{2}))'
                   ).format(R0, INV_END, R0 - 1, r=r))
    it.cell(r, 1).fill = PatternFill('solid', fgColor=AUTO)
    it.cell(r, 3).alignment = C('right')
    for col in (5, 6):
        it.cell(r, col).number_format = '0.000'
    it.cell(r, 6, '=IF(C{0}="","",N(D{0})*N(E{0}))'.format(r))
    it.cell(r, 6).fill = PatternFill('solid', fgColor=AUTO)

# ── لوحة متابعة الفواتير اليدوية (تتعبّى لحالها من ورقة «الفواتير») ──
# H رقم · I التاريخ (13/10) · J مجموع الفاتورة · K مجموع بنودها · L خلصت؟
for k, v in (('G', 3), ('H', 14), ('I', 13), ('J', 13), ('K', 13), ('L', 24)):
    it.column_dimensions[k].width = v
# فوق اللوحة: الفترة من أول تاريخ لآخر تاريخ (13/10) — والنموذج شهري: أكثر من شهر ← ⚠
it.merge_cells('H3:L3')
DMIN, DMAX = 'MIN($Q${0}:$Q${1})'.format(R0, INV_END), 'MAX($Q${0}:$Q${1})'.format(R0, INV_END)
DTXT = 'TEXT(DAY({0}),"00")&"/"&TEXT(MONTH({0}),"00")&"/"&YEAR({0})'
it['H3'] = ('=IF(COUNT($Q${0}:$Q${1})=0,"الفترة: —",'
            '"الفترة: من "&{a}&"  الى  "&{b}&'
            'IF(OR(MONTH({mn})<>MONTH({mx}),YEAR({mn})<>YEAR({mx})),"   ⚠ أكثر من شهر",""))'
            ).format(R0, INV_END, a=DTXT.format(DMIN), b=DTXT.format(DMAX), mn=DMIN, mx=DMAX)
it['H3'].fill = PatternFill('solid', fgColor=AUTO)
it['H3'].font = F(12, True, NAVY); it['H3'].alignment = C(); it['H3'].border = BOX
it.row_dimensions[3].height = 24
it.merge_cells('H4:L4')
it['H4'] = ('=IF(COUNTIF(\'الفواتير\'!$E${0}:$E${1},"يدوية")=0,"لا يوجد فواتير يدوية",'
            'COUNTIF(L{0}:L{1},"✅*")&" مكتملة من "&'
            'COUNTIF(\'الفواتير\'!$E${0}:$E${1},"يدوية")&" فاتورة يدوية")').format(R0, INV_END)
it['H4'].fill = PatternFill('solid', fgColor=AUTO)
it['H4'].font = F(12, True); it['H4'].alignment = C(); it['H4'].border = BOX
for i, h in enumerate(['رقم الفاتورة اليدوية', 'التاريخ', 'مجموع الفاتورة', 'مجموع بنودها', 'خلصت؟'], start=8):
    c = it.cell(6, i, h)
    c.font = F(11, True, 'FFFFFF'); c.alignment = C()
    c.fill = PatternFill('solid', fgColor=GOLD)
    c.border = Border(left=LINE, right=LINE, top=THICK, bottom=THICK)
SUMB = "SUMIF($B${0}:$B${1},$H{{r}},$F${0}:$F${1})".format(R0, ITEM_END)
CNTB = "COUNTIF($B${0}:$B${1},$H{{r}})".format(R0, ITEM_END)
for r in range(R0, INV_END + 1):
    it.cell(r, 8, ('=IF(\'الفواتير\'!$E{r}="يدوية",\'الفواتير\'!$D{r},"")').format(r=r))
    it.cell(r, 9, '=IF(OR($H{r}="",\'الفواتير\'!$C{r}=""),"",\'الفواتير\'!$C{r})'.format(r=r))
    it.cell(r, 10, '=IF($H{r}="","",\'الفواتير\'!$F{r})'.format(r=r))
    it.cell(r, 11, '=IF($H{r}="","",{s})'.format(r=r, s=SUMB.format(r=r)))
    it.cell(r, 12, ('=IF($H{r}="","",'
                    'IF($H{r}=0,"⛔ اكتب رقم الفاتورة",'
                    'IF({c}=0,"🟡 ما انزلت بنودها",'
                    'IF(ROUND($J{r}-$K{r},3)=0,"✅ خلصت",'
                    '"⛔ فرق "&TEXT($J{r}-$K{r},"0.000")))))').format(r=r, c=CNTB.format(r=r)))
    for col in range(8, 13):
        c = it.cell(r, col)
        c.border = BOX; c.alignment = C(); c.font = F(10, True) if col == 12 else F(11)
        c.fill = PatternFill('solid', fgColor=AUTO)
    it.cell(r, 9).number_format = '00"/"00"/"0000'
    for col in (10, 11):
        it.cell(r, col).number_format = '0.000'
    # Q (مخفي): تاريخ الفاتورة كتاريخ حقيقي — لكل الفواتير (مطبوعة ويدوية) — للفترة فوق
    it.cell(r, 17, ('=IF(OR(\'الفواتير\'!$B{r}="",\'الفواتير\'!$C{r}=""),"",IFERROR('
                    'DATE(MOD(\'الفواتير\'!$C{r},10000),MOD(INT(\'الفواتير\'!$C{r}/10000),100),INT(\'الفواتير\'!$C{r}/1000000)),""))').format(r=r))
LR = 'L{}:L{}'.format(R0, INV_END)
it.conditional_formatting.add(LR, FormulaRule(
    formula=['LEFT($L{},1)="✅"'.format(R0)], fill=PatternFill('solid', bgColor=GREEN), font=F(10, True, '17632A')))
it.conditional_formatting.add(LR, FormulaRule(
    formula=['LEFT($L{},1)="⛔"'.format(R0)], fill=PatternFill('solid', bgColor=RED), font=F(10, True, 'A11212')))
it.conditional_formatting.add(LR, FormulaRule(
    formula=['LEFT($L{},1)="🟡"'.format(R0)], fill=PatternFill('solid', bgColor='FFF0C2'), font=F(10, True, GOLD)))
it.conditional_formatting.add('H4', FormulaRule(
    formula=['AND(COUNTIF($L${0}:$L${1},"✅*")>0,COUNTIF($L${0}:$L${1},"✅*")='
             'COUNTIF(\'الفواتير\'!$E${0}:$E${1},"يدوية"))'.format(R0, INV_END)],
    fill=PatternFill('solid', bgColor=GREEN), font=F(12, True, '17632A')))
it.conditional_formatting.add('H3', FormulaRule(
    formula=['ISNUMBER(SEARCH("⚠",H3))'], fill=PatternFill('solid', bgColor='FFD9A0'), font=F(12, True, 'A11212')))

# ── القائمة المنسدلة: الفواتير اليدوية **اللي لسا ما خلصت** فقط ──
# N = ترتيب الفاتورة المفتوحة · O = القائمة مرصوصة فوق بلا فراغات · Q = التاريخ الحقيقي (كلهم مخفيين)
for k in ('M', 'N', 'O', 'P', 'Q'):
    it.column_dimensions[k].width = 14
    it.column_dimensions[k].hidden = True
it.cell(6, 14, 'ترتيب'); it.cell(6, 15, 'المفتوحة'); it.cell(6, 17, 'تاريخ')
for r in range(R0, INV_END + 1):
    it.cell(r, 14, ('=IF(AND($H{r}<>"",$H{r}<>0,LEFT($L{r},1)<>"✅"),'
                    'COUNT($N${h}:N{p})+1,"")').format(r=r, h=R0 - 1, p=r - 1))
    it.cell(r, 15, ('=IFERROR(INDEX($H${0}:$H${1},MATCH(ROW()-{2},$N${0}:$N${1},0)),"")'
                    ).format(R0, INV_END, R0 - 1))
# اسم معرَّف بطول متغيّر — عشان ما تطلع خانات فاضية بالقائمة
wb.defined_names.add(DefinedName('فواتير_مفتوحة', attr_text=(
    "OFFSET('البنود'!$O${0},0,0,MAX(1,SUMPRODUCT(--('البنود'!$O${0}:$O${1}<>\"\"))),1)"
).format(R0, INV_END)))
dv = DataValidation(type='list', formula1='=فواتير_مفتوحة',
                    allow_blank=True, showDropDown=False)
dv.error = 'اختر رقم فاتورة يدوية لسا ما خلصت من القائمة.'
dv.errorTitle = 'رقم غير موجود بالقائمة'
dv.showErrorMessage = False
it.add_data_validation(dv); dv.add('B{}:B{}'.format(R0, ITEM_END))

# بنود الفاتورة اللي خلصت ← باهتة (مغلقة) — أول قاعدة عشان تغلب (13/10)
it.conditional_formatting.add('A{}:F{}'.format(R0, ITEM_END), FormulaRule(
    formula=['AND($B{0}<>"",COUNTIFS($H${1}:$H${2},$B{0},$L${1}:$L${2},"✅*")>0)'.format(R0, R0, INV_END)],
    fill=PatternFill('solid', bgColor='F2F2F2'), font=Font(name='Arial', size=11, color='A6A6A6', italic=True),
    stopIfTrue=True))

it.conditional_formatting.add('E4', FormulaRule(
    formula=['LEFT(E4,1)="✅"'], fill=PatternFill('solid', bgColor=GREEN), font=F(12, True, '17632A')))
it.conditional_formatting.add('E4', FormulaRule(
    formula=['LEFT(E4,1)="⛔"'], fill=PatternFill('solid', bgColor=RED), font=F(12, True, 'A11212')))
it.conditional_formatting.add('A{}:F{}'.format(R0, ITEM_END), FormulaRule(
    formula=['$A{}="⛔"'.format(R0)], fill=PatternFill('solid', bgColor=RED)))
for col in ('B', 'D', 'E'):
    it.conditional_formatting.add('{0}{1}:{0}{2}'.format(col, R0, ITEM_END), FormulaRule(
        formula=['AND($C{0}<>"",${1}{0}="")'.format(R0, col)],
        fill=PatternFill('solid', bgColor=RED)))

W2 = ITEM_END + 2
it.merge_cells('A{0}:F{0}'.format(W2))
it['A{}'.format(W2)] = 'رقم الفاتورة لازم يكون نفسه المكتوب بورقة «الفواتير» — وإلا بيطلع ⛔ بعمود #'
it['A{}'.format(W2)].font = F(10, True, 'FFFFFF'); it['A{}'.format(W2)].alignment = C()
it['A{}'.format(W2)].fill = PatternFill('solid', fgColor=GOLD)
it.row_dimensions[W2].height = 22

page(it, 6, W2)
lock(it, 6, W2, [(c, R0, ITEM_END) for c in (2, 3, 4, 5)])

# ══════════ التعليمات ══════════
wi = wb.create_sheet('التعليمات')
wi.sheet_view.rightToLeft = True
wi.sheet_view.showGridLines = False
wi.column_dimensions['A'].width = 4
wi.column_dimensions['B'].width = 92
wi.merge_cells('A1:B1')
wi['A1'] = 'طريقة العمل'
wi['A1'].font = F(16, True, 'FFFFFF'); wi['A1'].alignment = C()
wi['A1'].fill = PatternFill('solid', fgColor=NAVY)
wi.row_dimensions[1].height = 30

STEPS = [
    ('sec', 'القاعدة الأساسية'),
    ('1', 'ملف واحد لكل شهر. بتفرّغ فيه فواتير شراء البضاعة **يوم بيوم** — مش آخر الشهر.'),
    ('2', 'اكتب الشهر مرة وحدة فوق بورقة «الفواتير» وبس. ورقة «البنود» بتاخده لحالها.'),
    ('sec', 'ورقة «الفواتير» — سطر لكل فاتورة'),
    ('3', 'المورد: اكتب الاسم **نفسه بالضبط** كل مرة. «كتانة» و«سليم كتانة» بينحسبوا موردين اثنين.'),
    ('4', 'التاريخ: اكتبه **أرقام ورا بعض بلا شرطات** — 24092026 والخانة بتعرضه 24/09/2026 لحالها. '
          'إذا صار **برتقالي** يعني التاريخ غلط (يوم أكبر من 31 أو شهر أكبر من 12).'),
    ('5', 'النوع: من القائمة — **مطبوعة** أو **يدوية**. هاي بتحدّد شو بعدها.'),
    ('6', '**مطبوعة** ← المورد + التاريخ + رقم الفاتورة + المجموع وبس. بنودها بتنقرأ من صورتها.'),
    ('6أ', 'صورة المطبوعة: صوّرها من فوق وبضوء منيح، وسمّيها **نفس عمود «اسم الصورة»** بالزبط (المورد_رقم الفاتورة.jpg)، '
           'وحطها بمجلّد «صور الفواتير» جوّا مجلّد الشهر.'),
    ('6ب', 'بعدها وقّف على خانة «الصورة» واكبس **Ctrl+K** واختار الصورة من المجلّد ← العمود الأخير بيصير **✓**. '
           'لحد ما ترفق الصورة الحالة بتضل «⛔ ارفق صورتها».'),
    ('7', '**يدوية** ← نفس الأربعة، وكمان بنودها بورقة «البنود». '
          'السبب: الخط اليدوي ما بينقرأ، والبنود لازم تنكتب عشان تنزل على النظام.'),
    ('8', 'الخصم: إذا في خصم على الفاتورة كلها اكتبه، وإلا اتركه فاضي. '
          'الصافي بيتحسب لحاله (المجموع − الخصم).'),
    ('9', 'رقم الفاتورة اللي بيصير **برتقالي** = مكرر بنفس الملف. غيّره، وإلا بنودها بتختلط.'),
    ('sec', 'ورقة «البنود» — للفواتير اليدوية فقط'),
    ('10', 'كل بند بسطر: رقم الفاتورة + البند + الكمية + الإفرادي. الإجمالي بيتحسب لحاله.'),
    ('11', 'رقم الفاتورة لازم يكون **نفسه** المكتوب بورقة «الفواتير». '
           'إذا طلع ⛔ مكان الرقم التسلسلي ← يعني الرقم مش موجود بورقة الفواتير.'),
    ('12', 'اكتب البند زي ما هو مكتوب بالفاتورة. مش لازم يطابق الكتالوج — أنا بطابقه.'),
    ('sec', 'لوحة المتابعة — يمين ورقة «البنود»'),
    ('◆', 'أرقام الفواتير **اليدوية** بتطلع فيها **لحالها** — ما بتكتب إشي. '
          'جنب كل رقم: مجموع الفاتورة · مجموع بنودها اللي أدخلتها · وعمود **«خلصت؟»**.'),
    ('◆', '🟡 **ما انزلت بنودها** = لسا ما بلّشت فيها · ⛔ **فرق 0.500** = المجموعان مش متطابقين '
          '· ✅ **خلصت** = المجموعان تطابقا وخلصت هاي الفاتورة.'),
    ('◆', 'فوق اللوحة بيطلع «**3 مكتملة من 5 فواتير يدوية**» — لما يصير أخضر يعني كلهن خلصوا.'),
    ('◆', 'جنب كل فاتورة **تاريخها**، وفوق اللوحة **الفترة من أول تاريخ لآخر تاريخ**. '
          'النموذج **شهري**: إذا طلع «⚠ أكثر من شهر» يعني في فاتورة من شهر ثاني.'),
    ('◆', 'الفاتورة اللي **خلصت** ← بنودها بتصير **باهتة (رمادي)** يعني مغلقة ومنتهية، وبتختفي من القائمة المنسدلة.'),
    ('sec', 'التدقيق — شغل الإكسل مش شغلك'),
    ('13', 'عمود «الحالة» بيقارن مجموع البنود مع مجموع الفاتورة لحاله: '
           '**✅ البنود مطابقة** يعني تمام · **⛔ فرق 0.500** يعني ناقص أو زايد 500 فلس.'),
    ('14', 'العدّاد فوق ما بيصير أخضر إلا لما كل فاتورة تكون سليمة. '
           'لا تبعث الملف وهو أحمر.'),
    ('sec', 'الإرسال'),
    ('15', 'الملف وصور الفواتير بمجلّد الشهر على الدرايف (مشتريات / شهر 10-2026) — ما في داعي تبعثهم.'),
    ('16', 'الفواتير بتنزل على أودو **يوم بيوم** — مش آخر الشهر. '
           'كل يوم بيمرّ بلا تنزيل = يوم بتشتغل فيه على أرقام مش صحيحة.'),
    ('sec', 'الفكرة باختصار'),
    ('◆', 'المطبوعة إلها صورة بتغني عن كتابة بنودها. اليدوية ما إلها — '
          'فبنكتب بنودها مرة وحدة هون، وبتضل مقروءة للأبد.'),
]
r = 3
for kind, txt in STEPS:
    if kind == 'sec':
        wi.merge_cells('A{0}:B{0}'.format(r))
        wi['A{}'.format(r)] = txt
        wi['A{}'.format(r)].font = F(12, True, 'FFFFFF')
        wi['A{}'.format(r)].alignment = C('right')
        wi['A{}'.format(r)].fill = PatternFill('solid', fgColor=TEAL)
        wi.row_dimensions[r].height = 24
    else:
        wi['A{}'.format(r)] = kind
        wi['A{}'.format(r)].font = F(11, True, NAVY)
        wi['A{}'.format(r)].alignment = C()
        wi['B{}'.format(r)] = txt.replace('**', '')
        wi['B{}'.format(r)].font = F(11)
        wi['B{}'.format(r)].alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
        wi.row_dimensions[r].height = 15 + 15 * (len(txt) // 78)
    r += 1
wi.print_area = 'A1:B{}'.format(r - 1)
wi.page_setup.paperSize = wi.PAPERSIZE_A4
wi.page_setup.fitToWidth = 1
wi.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print('تم:', OUT)
