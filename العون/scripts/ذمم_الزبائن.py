# -*- coding: utf-8 -*-
"""ذمم الزبائن — قائمة «مين عليه كم» للموظف (بأمر صاحب المحل 02/10/2026، بعد استشارة 4 خبراء)

  python3 scripts/ذمم_الزبائن.py            # يبني drive/ذمم_الزبائن.xlsx + .csv من أودو ويطبع الملخص
  python3 scripts/ذمم_الزبائن.py --اعرض     # يطبع الجدول بس (بلا كتابة)

القواعد:
  • المصدر الوحيد أودو: أسطر الذمم المدينة المرحّلة غير المسوّاة (account.move.line · asset_receivable · posted)
    مجمّعة بالشريك — بتشمل الدفعات غير المسوّاة (بالسالب) والمرتجعات، وما بتشمل المسودات.
  • الحسابات الشخصية (وسم «حساب شخصي») ما بتطلع — هاي لصاحب المحل بس.
  • الشيت للقراءة فقط — ممنوع أي كتابة فيه؛ أي دين أو دفعة بتنسجّل على أودو وبتطلع هون بالجولة الجاية.
  • الفرز من الأكبر للأصغر · عمر الدين من أقدم فاتورة مفتوحة · الحالة: 🟢 عادي · 🟡 فوق 30 يوم · 🔴 فوق 60 يوم = كاش فقط.
"""
import sys, os, csv, datetime as dt
from collections import defaultdict
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
sys.path.insert(0, os.path.join(ROOT, 'odoo_templates'))
from jrpc import x
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule

OUT_X = os.path.join(ROOT, 'drive', 'ذمم_الزبائن.xlsx')
OUT_C = os.path.join(ROOT, 'drive', 'ذمم_الزبائن.csv')
PERSONAL_TAG = 'حساب شخصي'
AMMAN = dt.timezone(dt.timedelta(hours=3))


def اجمع():
    """[(اسم, نوع, عليه, أقدم فاتورة, عمر, آخر دفعة تاريخ, آخر دفعة مبلغ, هاتف)] — من أودو."""
    tags = {t['id']: t['name'] for t in x('res.partner.category', 'search_read', [], ['id', 'name'])}
    lines = x('account.move.line', 'search_read',
              [('account_id.account_type', '=', 'asset_receivable'), ('parent_state', '=', 'posted'), ('reconciled', '=', False)],
              ['partner_id', 'amount_residual', 'date', 'move_id', 'debit', 'credit'])
    pay = x('account.move.line', 'search_read',
            [('account_id.account_type', '=', 'asset_receivable'), ('parent_state', '=', 'posted'), ('credit', '>', 0),
             ('move_id.move_type', 'not in', ['out_refund'])],
            ['partner_id', 'credit', 'date'], order='date desc')
    last = {}
    for l in pay:
        if l['partner_id'] and l['partner_id'][0] not in last: last[l['partner_id'][0]] = (l['date'], l['credit'])
    agg = defaultdict(lambda: {'عليه': 0.0, 'أقدم': None})
    for l in lines:
        if not l['partner_id']: continue
        a = agg[l['partner_id'][0]]; a['عليه'] += l['amount_residual']
        if l['debit'] > 0 and abs(l['amount_residual']) > 0.0005 and (a['أقدم'] is None or l['date'] < a['أقدم']): a['أقدم'] = l['date']
    ids = [p for p, a in agg.items() if abs(a['عليه']) > 0.0005]
    ps = {p['id']: p for p in x('res.partner', 'read', ids, ['name', 'phone', 'category_id'])} if ids else {}
    today = dt.datetime.now(AMMAN).date(); rows = []
    for pid in ids:
        p = ps[pid]
        if any(tags.get(c) == PERSONAL_TAG for c in p['category_id']): continue
        a = agg[pid]; d = dt.date.fromisoformat(a['أقدم']) if a['أقدم'] else None
        نوع = 'ورشة' if any(w in p['name'] for w in ('ورشة', 'مشروع', 'ذمم')) else 'بواقي'
        lp = last.get(pid)
        rows.append([p['name'], نوع, round(a['عليه'], 3), d, (today - d).days if d else None,
                     dt.date.fromisoformat(lp[0]) if lp else None, round(lp[1], 3) if lp else None, p['phone'] or ''])
    rows.sort(key=lambda r: -r[2])
    return rows


def ابنِ(rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'الذمم'
    ws.sheet_view.rightToLeft = True
    HEAD = ['#', 'اسم الزبون', 'النوع', 'عليه (دينار)', 'أقدم فاتورة مفتوحة', 'عمر الدين (يوم)', 'تاريخ آخر دفعة', 'آخر دفعة (دينار)', 'الهاتف', 'الحالة']
    for k, w in zip('ABCDEFGHIJ', [5, 30, 9, 14, 16, 13, 15, 14, 14, 14]): ws.column_dimensions[k].width = w
    F = lambda sz=11, b=False, c='000000': Font(name='Arial', size=sz, bold=b, color=c)
    thin = Side(style='thin', color='BFBFBF'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells('A1:J1'); ws['A1'] = 'ذمم الزبائن — محلات العون لمواد البناء'; ws['A1'].font = F(16, True, 'FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor='1F4E79'); ws['A1'].alignment = Alignment(horizontal='center', vertical='center'); ws.row_dimensions[1].height = 32
    ws['A2'] = 'آخر تحديث من أودو'; ws['B2'] = dt.datetime.now(AMMAN).replace(tzinfo=None, microsecond=0); ws['B2'].number_format = 'dd/mm/yyyy hh:mm'
    ws['A2'].font = F(10, True); ws['B2'].font = F(10, True)
    ws.merge_cells('D2:G2'); ws['D2'] = '=IF(NOW()-B2>4/24,"⚠️ البيانات قديمة — بلّغ صاحب المحل","✔ محدّثة")'
    ws['D2'].font = F(10, True); ws['D2'].alignment = Alignment(horizontal='center')
    ws['A3'] = 'المجموع عليهم'; ws['B3'] = '=SUM(D6:D1000)'; ws['B3'].number_format = '#,##0.000'
    ws['C3'] = 'عدد المدينين'; ws['D3'] = '=COUNTA(B6:B1000)'
    for c in ('A3', 'C3'): ws[c].font = F(10, True)
    for c in ('B3', 'D3'): ws[c].font = F(11, True, '1F4E79')
    ws.merge_cells('F3:J3'); ws['F3'] = 'الشيت للقراءة فقط — أي دين أو دفعة بتنسجّل على أودو وبتطلع هون لحالها'
    ws['F3'].font = F(9, False, '7F7F7F'); ws['F3'].alignment = Alignment(horizontal='right')
    ws.merge_cells('A4:J4'); ws['A4'] = '🟢 عادي (لحد 30 يوم)   ·   🟡 فوق 30 يوم: ذكّره بلطف   ·   🔴 فوق 60 يوم: كاش فقط لحد ما يسدّ'
    ws['A4'].font = F(10, True, '404040'); ws['A4'].alignment = Alignment(horizontal='center')
    for i, h in enumerate(HEAD, start=1):
        c = ws.cell(5, i, h); c.font = F(10, True); c.fill = PatternFill('solid', fgColor='BDD7EE')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = box
    ws.row_dimensions[5].height = 30; ws.freeze_panes = 'A6'
    for r, row in enumerate(rows, start=6):
        name, نوع, amt, old, age, lpd, lpa, phone = row
        vals = [r - 5, name, نوع, amt, old, None, lpd, lpa, phone, None]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(r, i, v); c.border = box; c.font = F(11, i == 4); c.alignment = Alignment(horizontal='right' if i == 2 else 'center', vertical='center')
        ws.cell(r, 4).number_format = '#,##0.000'; ws.cell(r, 8).number_format = '#,##0.000'
        ws.cell(r, 5).number_format = 'dd/mm/yyyy'; ws.cell(r, 7).number_format = 'dd/mm/yyyy'
        ws.cell(r, 6, '=IF(E%d="","",TODAY()-E%d)' % (r, r))
        ws.cell(r, 10, '=IF(F%d="","",IF(F%d>60,"🔴 كاش فقط",IF(F%d>30,"🟡 ذكّره","🟢 عادي")))' % (r, r, r))
        ws.row_dimensions[r].height = 22
    end = 5 + max(len(rows), 1)
    rng = 'A6:J%d' % end
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$F6>60'], fill=PatternFill('solid', fgColor='FFC7CE')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['AND($F6>30,$F6<=60)'], fill=PatternFill('solid', fgColor='FFEB9C')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$D6<0'], font=Font(color='1E5631', bold=True)))
    wb.save(OUT_X)
    with open(OUT_C, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(HEAD[1:9])
        for row in rows: w.writerow([row[0], row[1], '%.3f' % row[2], row[3] or '', row[4] if row[4] is not None else '', row[5] or '', ('%.3f' % row[6]) if row[6] else '', row[7]])


def main():
    rows = اجمع()
    tot = sum(r[2] for r in rows)
    print('ذمم الزبائن — %d زبون · المجموع %.3f · %s' % (len(rows), tot, dt.datetime.now(AMMAN).strftime('%d/%m/%Y %H:%M')))
    for r in rows:
        print('  %-30s %-6s %9.3f  أقدم %s  عمر %s  آخر دفعة %s %s  %s' % (r[0][:30], r[1], r[2], r[3].strftime('%d/%m') if r[3] else '—',
              r[4] if r[4] is not None else '—', r[5].strftime('%d/%m') if r[5] else '—', ('%.3f' % r[6]) if r[6] else '', r[7]))
    if '--اعرض' not in sys.argv:
        ابنِ(rows); print('✔', OUT_X)


if __name__ == '__main__':
    main()
