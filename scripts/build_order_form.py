# -*- coding: utf-8 -*-
"""نموذج طلبية ورشة — محلات العون لمواد البناء
نموذج موحّد لكل الورشات، بضوابط تدقيق مدمجة تمنع سهو البنود.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = '/home/user/Fam-Bam/نموذج_طلبية_ورشة.xlsx'

R0, R1 = 12, 46            # أول وآخر صف بيانات
NROWS = R1 - R0 + 1

NAVY   = '1F4E6B'
TEAL   = '0F6E6E'
EDIT   = 'FFF6CC'          # أصفر = خانة إدخال
AUTO   = 'EDEDED'          # رمادي = محسوبة تلقائياً — لا تُعدّل
GREEN  = 'D6F0D6'
RED    = 'FBD5D5'
LINE   = Side(style='thin', color='9AA5AD')
THICK  = Side(style='medium', color=NAVY)
BOX    = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)

def F(sz=11, b=False, c='000000'):
    return Font(name='Arial', size=sz, bold=b, color=c)
def C(h='center'):
    return Alignment(horizontal=h, vertical='center', wrap_text=True)

wb = openpyxl.Workbook()

# ══════════════════════ ورقة الطلبية ══════════════════════
ws = wb.active
ws.title = 'طلبية'
ws.sheet_view.rightToLeft = True
ws.sheet_view.showGridLines = False

WID = {'A': 5, 'B': 40, 'C': 8.5, 'D': 10, 'E': 12, 'F': 12, 'G': 12, 'H': 8, 'I': 22}
for k, v in WID.items():
    ws.column_dimensions[k].width = v

# ── العنوان ──
ws.merge_cells('A1:I1')
ws['A1'] = 'نموذج طلبية ورشة — محلات العون لمواد البناء'
ws['A1'].font = F(16, True, 'FFFFFF'); ws['A1'].alignment = C()
ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
ws.row_dimensions[1].height = 30

ws.merge_cells('A2:I2')
ws['A2'] = 'تُملأ من ورقة الطلبية الأصلية / رسالة الواتساب — والأصل يُحفظ بنفس رقم الطلبية'
ws['A2'].font = F(9, False, '555555'); ws['A2'].alignment = C()
ws.row_dimensions[2].height = 16

# ── ترويسة البيانات ──
def lab(cell, txt):
    ws[cell] = txt; ws[cell].font = F(10, True); ws[cell].alignment = C('right')
    ws[cell].fill = PatternFill('solid', fgColor='E8EDF0'); ws[cell].border = BOX
def cells(sheet, rng):
    v = sheet[rng]
    if hasattr(v, 'value'):
        return [v]
    out = []
    for row in v:
        if hasattr(row, 'value'):
            out.append(row)
        else:
            out.extend(row)
    return out

def inp(rng):
    first = rng.split(':')[0]
    if ':' in rng: ws.merge_cells(rng)
    ws[first].fill = PatternFill('solid', fgColor=EDIT)
    ws[first].font = F(11, True); ws[first].alignment = C()
    for c in cells(ws, rng): c.border = BOX

for r in (3, 4):
    ws.row_dimensions[r].height = 22
lab('A3', 'اسم الورشة / المشروع'); inp('B3:D3')
lab('E3', 'رقم الطلبية');          inp('F3')
lab('G3', 'تاريخ الطلبية');        inp('H3:I3')
lab('A4', 'المهندس / المستلِم');   inp('B4:D4')
lab('E4', 'مصدر الطلبية');         inp('F4')
lab('G4', 'صورة الأصل محفوظة؟');   inp('H4:I4')

# الطلبية الطويلة تنقسم ورقتين — كل ورقة تحمل ضوابطها هي
ws.row_dimensions[5].height = 22
lab('A5', 'الورقة رقم');  inp('B5')
lab('C5', 'من أصل');      inp('D5')
ws.merge_cells('E5:I5')
ws['E5'] = 'الطلبية الطويلة تنقسم ورقتين — وكل ورقة بتحمل عدد بنودها هي، لا عدد الطلبية كلها'
ws['E5'].font = F(9, False, '555555'); ws['E5'].alignment = C()
ws['E5'].border = BOX

# ── شريط ضوابط التدقيق ──
ws.merge_cells('A6:I6')
ws['A6'] = '◆ ضوابط التدقيق — تُملأ من الورقة الأصلية قبل البدء بالإدخال ◆'
ws['A6'].font = F(11, True, 'FFFFFF'); ws['A6'].alignment = C()
ws['A6'].fill = PatternFill('solid', fgColor=TEAL)
ws.row_dimensions[6].height = 24

def auto(cell):
    ws[cell].fill = PatternFill('solid', fgColor=AUTO)
    ws[cell].font = F(11, True); ws[cell].alignment = C(); ws[cell].border = BOX

for r in (7, 8, 9):
    ws.row_dimensions[r].height = 24

lab('A7', 'عدد البنود حسب الأصل'); inp('B7')
lab('C7', 'المُدخَل فعلياً')
ws['D7'] = '=COUNTA(B{}:B{})'.format(R0, R1); auto('D7')
ws.merge_cells('E7:I7'); auto('E7')
ws['E7'] = ('=IF(B7="","— اكتب عدد بنود هاي الورقة أولاً",'
            'IF(D7=B7,"✅ مطابق — ما في بند ناقص",'
            'IF(D7<B7,"⛔ ناقص "&(B7-D7)&" بند — راجع الورقة",'
            '"⛔ زائد "&(D7-B7)&" بند — في بند مكرر")))')

lab('A8', 'مجموع الكميات حسب الأصل'); inp('B8')
lab('C8', 'المُدخَل فعلياً')
ws['D8'] = '=SUM(C{}:C{})'.format(R0, R1); auto('D8')
ws.merge_cells('E8:I8'); auto('E8')
ws['E8'] = ('=IF(B8="","— اختياري (موصى به للطلبيات فوق 10 بنود)",'
            'IF(ROUND(D8,3)=ROUND(B8,3),"✅ مطابق",'
            '"⛔ فرق "&TEXT(D8-B8,"0.###")&" — في كمية مكتوبة غلط"))')

lab('A9', 'البنود المدقَّقة ✔')
ws['B9'] = '=COUNTIF(H{}:H{},"✔")'.format(R0, R1); auto('B9')
lab('C9', 'الحالة النهائية')
ws.merge_cells('D9:I9'); auto('D9')
ws['D9'] = ('=IF(D7=0,"⛔ ما في بنود مُدخلة",'
            'IF(B7="","⛔ ناقص: عدد البنود حسب الأصل",'
            'IF(D7<>B7,"⛔ عدد البنود مش مطابق",'
            'IF(B9<D7,"⛔ باقي "&(D7-B9)&" بند بدون تدقيق",'
            '"✅ جاهز للترحيل — ابعثها للأستاذ"))))')
ws['D9'].font = F(12, True)

# ── رأس الجدول ──
HDR = ['#', 'البند (كما طلع من المحل)', 'الكمية', 'الوحدة',
       'السعر الإفرادي', 'المجموع', 'تاريخ الإخراج', '✔ تدقيق', 'ملاحظات']
ws.row_dimensions[11].height = 34
for i, h in enumerate(HDR, start=1):
    c = ws.cell(11, i, h)
    c.font = F(10, True, 'FFFFFF'); c.alignment = C()
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.border = Border(left=LINE, right=LINE, top=THICK, bottom=THICK)

# ── صفوف البيانات ──
for r in range(R0, R1 + 1):
    ws.row_dimensions[r].height = 19
    ws.cell(r, 1, '=IF(B{0}="","",ROW()-{1})'.format(r, R0 - 1))
    ws.cell(r, 6, '=IF(OR(C{0}="",E{0}=""),"",C{0}*E{0})'.format(r))
    for col in range(1, 10):
        c = ws.cell(r, col)
        c.border = BOX; c.font = F(11)
        c.alignment = C('right' if col in (2, 9) else 'center')
        if col in (1, 6):                       # محسوبة
            c.fill = PatternFill('solid', fgColor=AUTO)
        else:
            c.fill = PatternFill('solid', fgColor='FFFFFF')
    ws.cell(r, 3).number_format = '0.###'
    ws.cell(r, 5).number_format = '0.000'
    ws.cell(r, 6).number_format = '0.000'
    ws.cell(r, 7).number_format = 'DD/MM/YYYY'

# ── المجموع ──
TR = R1 + 1
ws.row_dimensions[TR].height = 26
ws.merge_cells('A{0}:E{0}'.format(TR))
ws['A{}'.format(TR)] = 'المجموع الكلي للبنود المسعَّرة (د.أ)'
ws['A{}'.format(TR)].font = F(11, True); ws['A{}'.format(TR)].alignment = C('right')
ws['A{}'.format(TR)].fill = PatternFill('solid', fgColor='E8EDF0')
ws['F{}'.format(TR)] = '=SUM(F{}:F{})'.format(R0, R1)
ws['F{}'.format(TR)].font = F(12, True); ws['F{}'.format(TR)].alignment = C()
ws['F{}'.format(TR)].number_format = '0.000'
ws['F{}'.format(TR)].fill = PatternFill('solid', fgColor=AUTO)
ws.merge_cells('G{0}:I{0}'.format(TR))
ws['G{}'.format(TR)] = 'البنود بدون سعر تُسعَّر من الكتالوج'
ws['G{}'.format(TR)].font = F(9, False, '555555'); ws['G{}'.format(TR)].alignment = C()
for col in range(1, 10):
    ws.cell(TR, col).border = Border(left=LINE, right=LINE, top=THICK, bottom=THICK)

# ── التواقيع ──
SR = TR + 2
ws.merge_cells('A{0}:I{0}'.format(SR))
ws['A{}'.format(SR)] = 'لا تُرسَل الطلبية إلا والحالة النهائية خضراء ✅'
ws['A{}'.format(SR)].font = F(10, True, 'FFFFFF'); ws['A{}'.format(SR)].alignment = C()
ws['A{}'.format(SR)].fill = PatternFill('solid', fgColor=TEAL)
ws.row_dimensions[SR].height = 22

for j, (t1, t2) in enumerate([('أدخلها (الموظف)', 'قرأ الورقة وأدخلها بند بند'),
                              ('دقّقها (أبو علي)', 'قارَنها بالورقة الأصلية بنداً وكميةً')]):
    r = SR + 1 + j
    ws.row_dimensions[r].height = 26
    lab('A{}'.format(r), t1); inp('B{}:C{}'.format(r, r))
    lab('D{}'.format(r), 'التاريخ'); inp('E{}'.format(r))
    lab('F{}'.format(r), 'التوقيع'); inp('G{}'.format(r))
    ws.merge_cells('H{0}:I{0}'.format(r))
    ws['H{}'.format(r)] = t2
    ws['H{}'.format(r)].font = F(8, False, '555555'); ws['H{}'.format(r)].alignment = C()
    ws['H{}'.format(r)].border = BOX

# ── قوائم منسدلة ──
def dv(formula, cells, msg=None):
    d = DataValidation(type='list', formula1=formula, allow_blank=True, showDropDown=False)
    ws.add_data_validation(d)
    d.add(cells)
    return d

dv('"حبة,ربطة,جوز,لفة,باكيت,شوال,درم,جالون,طقم,عبوة,متر,كغم,لتر,طن"',
   'D{}:D{}'.format(R0, R1))
dv('"✔"', 'H{}:H{}'.format(R0, R1))
dv('"ورقة مكتوبة,واتساب,هاتف"', 'F4')
dv('"نعم,لا"', 'H4')

num = DataValidation(type='decimal', operator='greaterThan', formula1='0',
                     allow_blank=True, showErrorMessage=True,
                     error='الكمية لازم تكون رقم فقط. الوحدة (ربطة/جوز/متر) تُكتب بخانة الوحدة.',
                     errorTitle='كمية غير صحيحة')
ws.add_data_validation(num); num.add('C{}:C{}'.format(R0, R1))

# ── تنسيق شرطي ──
rng_status = ['E7', 'E8', 'D9']
for cell in rng_status:
    ws.conditional_formatting.add(cell, FormulaRule(
        formula=['LEFT({},1)="✅"'.format(cell)],
        fill=PatternFill('solid', bgColor=GREEN), font=F(11, True, '17632A')))
    ws.conditional_formatting.add(cell, FormulaRule(
        formula=['LEFT({},1)="⛔"'.format(cell)],
        fill=PatternFill('solid', bgColor=RED), font=F(11, True, 'A11212')))

body = 'A{}:I{}'.format(R0, R1)
# بند بدون كمية
ws.conditional_formatting.add('C{}:C{}'.format(R0, R1), FormulaRule(
    formula=['AND($B{}<>"",$C{}="")'.format(R0, R0)],
    fill=PatternFill('solid', bgColor=RED)))
# بند بدون تاريخ
ws.conditional_formatting.add('G{}:G{}'.format(R0, R1), FormulaRule(
    formula=['AND($B{}<>"",$G{}="")'.format(R0, R0)],
    fill=PatternFill('solid', bgColor=RED)))
# سطر مدقَّق = أخضر خفيف
ws.conditional_formatting.add(body, FormulaRule(
    formula=['$H{}="✔"'.format(R0)],
    fill=PatternFill('solid', bgColor='F0F8F0')))

# ── الطباعة ──
ws.print_area = 'A1:I{}'.format(SR + 2)
ws.page_setup.orientation = 'landscape'
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.4
ws.print_title_rows = '11:11'
ws.freeze_panes = 'A12'

# ── حماية: الخانات المحسوبة فقط مقفلة ──
from openpyxl.styles import Protection
for row in ws.iter_rows(min_row=1, max_row=SR + 2, max_col=9):
    for c in row:
        c.protection = Protection(locked=True)
for rng in ['B3:D3', 'F3', 'H3:I3', 'B4:D4', 'F4', 'H4:I4', 'B5', 'D5', 'B7', 'B8',
            'B{}:E{}'.format(R0, R1), 'G{}:I{}'.format(R0, R1),
            'B{}:C{}'.format(SR + 1, SR + 1), 'E{}'.format(SR + 1), 'G{}'.format(SR + 1),
            'B{}:C{}'.format(SR + 2, SR + 2), 'E{}'.format(SR + 2), 'G{}'.format(SR + 2)]:
    for c in cells(ws, rng):
        c.protection = Protection(locked=False)
ws.protection.sheet = True
ws.protection.formatCells = False
ws.protection.selectLockedCells = False

# ══════════════════════ ورقة التعليمات ══════════════════════
wi = wb.create_sheet('التعليمات')
wi.sheet_view.rightToLeft = True
wi.sheet_view.showGridLines = False
wi.column_dimensions['A'].width = 4
wi.column_dimensions['B'].width = 95

wi.merge_cells('A1:B1')
wi['A1'] = 'طريقة العمل — نموذج طلبية الورشة'
wi['A1'].font = F(16, True, 'FFFFFF'); wi['A1'].alignment = C()
wi['A1'].fill = PatternFill('solid', fgColor=NAVY)
wi.row_dimensions[1].height = 32

STEPS = [
    ('sec', 'أولاً — الموظف اللي بطلّع البضاعة'),
    ('1', 'كل طلبية إلها رقم متسلسل. اكتب الرقم **على الورقة نفسها** قبل أي شي، وافتح ورقة جديدة بالإكسل بنفس الرقم.'),
    ('2', 'قبل ما تبلّش الإدخال: **عُدّ بنود الورقة** واكتب العدد بخانة «عدد البنود حسب الأصل». هاي أهم خانة بالنموذج كله.'),
    ('3', 'أدخل البنود بند بند. وكل بند بتدخله — **اشطبه بالقلم على الورقة**. أي بند بضل بدون شطب = بند نسيته.'),
    ('4', 'الكمية رقم فقط. «2 ربطة» تُكتب: الكمية 2 والوحدة «ربطة». لا تكتب الوحدة جوا خانة الكمية.'),
    ('5', 'تاريخ الإخراج إلزامي لكل بند — كل تاريخ بصير فاتورة لحاله بالنظام.'),
    ('6', 'الطلبية الطويلة (فوق 35 بند) تنقسم ورقتين بنفس رقم الطلبية: «الورقة 1 من 2» و«الورقة 2 من 2». '
          'كل ورقة بتحمل عدد بنودها هي — مثلاً 35 بالأولى و15 بالثانية — مش عدد الطلبية كلها.'),
    ('7', 'السعر الإفرادي: إذا المهندس اتفق مع المقاول على سعر، اكتبه هون مباشرة. إذا تركته فاضي بنسعّره من الكتالوج.'),
    ('8', 'صوّر الورقة الأصلية واحفظها بنفس رقم الطلبية، وحطّ «نعم» بخانة «صورة الأصل محفوظة».'),
    ('sec', 'ثانياً — التدقيق (أبو علي)'),
    ('9', 'التدقيق بصير **من الورقة الأصلية، مش من الشاشة**. امسك الورقة واقرأ منها، وقارن كل سطر بالإكسل.'),
    ('10', 'لكل بند تتأكد منه (الاسم صح + الكمية صح) حُطّ ✔ بخانة التدقيق.'),
    ('11', 'العدّاد فوق بيعدّ لحالو. ما بتصير الحالة خضراء إلا لما كل بند يتدقّق وعدد البنود يطابق الورقة.'),
    ('12', 'وقّع بخانة «دقّقها» وبعدين ابعث الملف + صورة الورقة للأستاذ.'),
    ('sec', 'الضوابط المدمجة بالنموذج'),
    ('•', '«عدد البنود حسب الأصل» ≠ عدد البنود المُدخلة ← بطلع ⛔ أحمر ويقلك كم بند ناقص.'),
    ('•', '«مجموع الكميات» (اختياري، موصى به فوق 10 بنود) ← بيمسك الكمية المكتوبة غلط (5 بدل 50).'),
    ('•', 'بند مكتوب بدون كمية أو بدون تاريخ ← الخانة بتصير حمرا.'),
    ('•', 'الخانات الرمادية محسوبة تلقائياً ومقفلة — لا تحاول تعدّلها.'),
    ('sec', 'الفكرة باختصار'),
    ('◆', 'البند المنسي ما بيترك أثر على الإكسل — عشان هيك التدقيق لازم يكون مقابل الورقة الأصلية، '
          'وعشان هيك في شطب بالقلم على الورقة + عدّ مسبق للبنود + شخص تاني بقرأ من الأصل. '
          'كل ضابط منهم بمسك الشي اللي الثاني بفوّته.'),
]

r = 3
for kind, txt in STEPS:
    if kind == 'sec':
        wi.merge_cells('A{0}:B{0}'.format(r))
        wi['A{}'.format(r)] = txt
        wi['A{}'.format(r)].font = F(12, True, 'FFFFFF')
        wi['A{}'.format(r)].fill = PatternFill('solid', fgColor=TEAL)
        wi['A{}'.format(r)].alignment = C()
        wi.row_dimensions[r].height = 24
        r += 1
        continue
    wi['A{}'.format(r)] = kind
    wi['A{}'.format(r)].font = F(11, True, NAVY); wi['A{}'.format(r)].alignment = C()
    wi['B{}'.format(r)] = txt.replace('**', '')
    wi['B{}'.format(r)].font = F(11); wi['B{}'.format(r)].alignment = C('right')
    wi.row_dimensions[r].height = 30 if len(txt) < 110 else 46
    r += 1

wi.print_area = 'A1:B{}'.format(r - 1)
wi.page_setup.orientation = 'portrait'
wi.page_setup.paperSize = wi.PAPERSIZE_A4
wi.page_setup.fitToWidth = 1
wi.page_setup.fitToHeight = 1
wi.sheet_properties.pageSetUpPr.fitToPage = True

wb.active = 0
wb.save(OUT)
print('تم:', OUT)
