# -*- coding: utf-8 -*-
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SP='/tmp/claude-0/-home-user-Fam-Bam/a76a5bbe-02f9-54b9-9aba-c7781c667b82/scratchpad'
data = json.load(open(SP+'/yz_export.json'))

F  = 'Arial'
HDR  = PatternFill('solid', fgColor='4A5D73')
EDIT = PatternFill('solid', fgColor='FFF2CC')   # أصفر = للتعبئة
INVF = PatternFill('solid', fgColor='EEF2F7')
EXF  = PatternFill('solid', fgColor='EAF3EA')
thin = Side(style='thin', color='AAAAAA')
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)
H    = Font(name=F, bold=True, color='FFFFFF', size=11)
B    = Font(name=F, bold=True, size=11)
N    = Font(name=F, size=10)
G    = Font(name=F, size=9, color='777777')
C    = Alignment(horizontal='center', vertical='center')
R    = Alignment(horizontal='right', vertical='center', wrap_text=True)

wb = Workbook()

# ───────────────────────── ورقة ١: التعليمات
ws = wb.active; ws.title='التعليمات'; ws.sheet_view.rightToLeft=True
rows = [
 ('كشف تدقيق فواتير يزيد العمري', None),
 ('محلات العون لمواد البناء — 18/09/2026', None),
 (None,None),
 ('الخانات الصفراء هي الوحيدة الي بتعبّيها. باقي الأعمدة للقراءة بس.', None),
 (None,None),
 ('ورقة «البنود» — سطر لكل بند', None),
 ('التاريخ الصحيح', 'اكتب التاريخ الصح لو تاريخ الفاتورة غلط (يوم/شهر/سنة). اتركه فاضي لو التاريخ مضبوط.'),
 ('الكمية الصحيحة', 'اكتب الكمية الصح لو مختلفة. اتركها فاضية لو مضبوطة.'),
 ('السعر الصحيح', 'اكتب سعر الحبة الصح لو مختلف. اتركه فاضي لو مضبوط.'),
 ('احذف؟', 'اكتب «نعم» لو البند مش موجود بالفاتورة الأصلية ولازم ينشال.'),
 ('مدقّق؟', 'اكتب «نعم» لما تتأكد إن البند صحيح ١٠٠٪.'),
 ('ملاحظات', 'أي إشي بدك تقوله عن البند — اسم الصنف غلط، بند ناقص، إلخ.'),
 (None,None),
 ('ورقة «الفواتير» — سطر لكل فاتورة', None),
 ('التاريخ الصحيح', 'لو بدك تصحّح تاريخ الفاتورة كلها من مكان واحد.'),
 ('مدقّقة؟', 'اكتب «نعم» لما تخلّص كل بنود الفاتورة.'),
 ('بنود ناقصة', 'اكتب البنود الي بالورقة ومش بأودو: الاسم × الكمية @ السعر'),
 ('ملاحظات', 'أي إشي عن الفاتورة ككل.'),
 (None,None),
 ('مهم', None),
 ('لا تغيّر عمود «المعرّف» — هو الي بربط فيه البند بأودو.', None),
 ('عمود «الإجمالي الجديد» بينحسب لحاله لما تكتب كمية أو سعر.', None),
 ('لما تخلّص ابعتلي الملف وبنفّذ كل التعديلات دفعة وحدة.', None),
]
for i,(a,b) in enumerate(rows, start=1):
    if a: ws.cell(i,1,a).font = B if i in (1,6,14,20) else N
    if b: ws.cell(i,2,b).font = N; ws.cell(i,2).alignment=R
ws.cell(1,1).font = Font(name=F, bold=True, size=14)
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 90

# ───────────────────────── ورقة ٢: الفواتير
ws = wb.create_sheet('الفواتير'); ws.sheet_view.rightToLeft=True
cols = [('المعرّف',10),('رقم الفاتورة',20),('النوع',12),('التاريخ بأودو',15),('عدد البنود',11),
        ('المجموع بأودو',15),('التاريخ الصحيح',16),('مدقّقة؟',11),('بنود ناقصة',42),('ملاحظات',34)]
for j,(t,w) in enumerate(cols, start=1):
    c=ws.cell(1,j,t); c.font=H; c.fill=HDR; c.alignment=C; c.border=BOX
    ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes='A2'
r=2
for m in data:
    vals=[m['id'], m['name'], 'مرتجع' if m['move_type']=='out_refund' else 'فاتورة',
          m['invoice_date'], len([l for l in m['lines'] if l['display_type']=='product']),
          m['amount_total'], None, None, None, None]
    for j,v in enumerate(vals, start=1):
        c=ws.cell(r,j,v); c.font=N; c.border=BOX
        c.alignment = C if j in (1,3,4,5,6,7,8) else R
        if j==6: c.number_format='#,##0.000'
        if j>=7: c.fill=EDIT
    r+=1
last=r-1
ws.cell(r,5,'المجموع').font=B; ws.cell(r,5).alignment=C; ws.cell(r,5).border=BOX
c=ws.cell(r,6,'=SUM(F2:F%d)'%last); c.font=B; c.number_format='#,##0.000'; c.alignment=C; c.border=BOX
dv=DataValidation(type='list', formula1='"نعم"', allow_blank=True); ws.add_data_validation(dv)
dv.add('H2:H%d'%last)

# ───────────────────────── ورقة ٣: البنود
ws = wb.create_sheet('البنود'); ws.sheet_view.rightToLeft=True
cols = [('المعرّف',10),('رقم الفاتورة',18),('التاريخ بأودو',14),('#',5),('كود الصنف',13),
        ('اسم الصنف / البيان',46),('الكمية',9),('السعر',10),('الإجمالي',12),
        ('التاريخ الصحيح',15),('الكمية الصحيحة',14),('السعر الصحيح',13),('الإجمالي الجديد',15),
        ('احذف؟',9),('مدقّق؟',9),('ملاحظات',34)]
for j,(t,w) in enumerate(cols, start=1):
    c=ws.cell(1,j,t); c.font=H; c.fill=HDR; c.alignment=C; c.border=BOX
    ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes='E2'
ws.row_dimensions[1].height=30

# سطر مثال
ex=['(مثال)','INV/2024/00001','2024-04-13',1,'100114','ريشة حجر عادي 08 ملم',10,0.375,3.75,
    '2024-04-15',12,0.400,'=IFERROR(IF(K2="",G2,K2)*IF(L2="",H2,L2),"")','','نعم','السعر بالورقة 0.400 مش 0.375']
for j,v in enumerate(ex, start=1):
    c=ws.cell(2,j,v); c.font=G; c.fill=EXF; c.border=BOX
    c.alignment = C if j in (1,3,4,5,7,8,9,10,11,12,13,14,15) else R
    if j in (8,9,12,13): c.number_format='#,##0.000'

r=3
for m in data:
    n=0
    for l in m['lines']:
        if l['display_type']=='product': n+=1
        code = (l['product_id'][1].split(']')[0].lstrip('[') if l['product_id'] and ']' in l['product_id'][1] else '')
        nm = ' '.join((l['name'] or '').split())
        if l['display_type']!='product':
            nm = ('— قسم — ' if l['display_type']=='line_section' else '— ملاحظة — ') + nm
        vals=[l['id'], m['name'], m['invoice_date'], n if l['display_type']=='product' else '',
              code, nm,
              l['quantity'] if l['display_type']=='product' else '',
              l['price_unit'] if l['display_type']=='product' else '',
              l['price_subtotal'] if l['display_type']=='product' else '',
              None,None,None,
              '=IFERROR(IF(K{0}="",G{0},K{0})*IF(L{0}="",H{0},L{0}),"")'.format(r) if l['display_type']=='product' else None,
              None,None,None]
        for j,v in enumerate(vals, start=1):
            c=ws.cell(r,j,v); c.font=N; c.border=BOX
            c.alignment = C if j in (1,3,4,5,7,8,9,10,11,12,13,14,15) else R
            if j in (8,9,12,13): c.number_format='#,##0.000'
            if j>=10 and j!=13: c.fill=EDIT
            if j==13: c.fill=INVF
        r+=1
last=r-1
ws.cell(r,6,'المجموع').font=B; ws.cell(r,6).alignment=C; ws.cell(r,6).border=BOX
for col,let in ((9,'I'),(13,'M')):
    c=ws.cell(r,col,'=SUM(%s3:%s%d)'%(let,let,last)); c.font=B
    c.number_format='#,##0.000'; c.alignment=C; c.border=BOX
dv2=DataValidation(type='list', formula1='"نعم"', allow_blank=True); ws.add_data_validation(dv2)
dv2.add('N3:N%d'%last); dv2.add('O3:O%d'%last)

wb.save('/home/user/Fam-Bam/تدقيق_فواتير_يزيد_العمري.xlsx')
print('✔ انعمل · فواتير %d · سطور %d' % (len(data), last-2))
