# -*- coding: utf-8 -*-
"""Workbook for entering a project's back-log of handwritten invoices.

The shop wrote weeks of tickets for one workshop on paper. This builds the
sheet those tickets get copied into, with the live catalogue beside it so a
line either resolves to a real product or is flagged for a decision before
anything reaches Odoo.

Three things a line can be, and the sheet says which:
  • matched to a priced product          -> ready
  • matched to a product priced at zero  -> needs a price
  • matched to nothing                   -> new product, or a name to fix

    python scripts/build_project_invoices_workbook.py "مشروع المدرسة" [out.xlsx]
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from scripts.check_connection import load_dotenv

load_dotenv(Path(__file__).resolve().parent.parent / ".env")

from odoo_mcp.client import OdooClient                      # noqa: E402
from odoo_mcp.config import OdooConfig                      # noqa: E402
from openpyxl import Workbook                               # noqa: E402
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # noqa: E402
from openpyxl.utils import get_column_letter                # noqa: E402
from openpyxl.comments import Comment                       # noqa: E402

ROWS = 400

INK, SOFT = "14171C", "6B7480"
HEAD_BG, HEAD_FG = "0F5C63", "FFFFFF"
TYPE_BG, AUTO_BG, NOTE_BG = "FFF6D9", "E9EFF0", "F4F6F7"

_thin = Side(style="thin", color="C6CCD4")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

# One object per style, reused — a fresh Font per cell bloats the file and
# makes both openpyxl and the reader crawl on a 4500-row sheet.
F_HEAD = Font(name="Arial", bold=True, size=10, color=HEAD_FG)
F_BODY = Font(name="Arial", size=10, color=INK)
F_SMALL = Font(name="Arial", size=9, color=INK)
F_AUTO = Font(name="Arial", size=10, color=SOFT)
F_BOLD = Font(name="Arial", bold=True, size=10, color=INK)
FILL_HEAD = PatternFill("solid", fgColor=HEAD_BG)
FILL_TYPE = PatternFill("solid", fgColor=TYPE_BG)
FILL_AUTO = PatternFill("solid", fgColor=AUTO_BG)
FILL_NOTE = PatternFill("solid", fgColor=NOTE_BG)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

COLS = [
    ("رقم الفاتورة",        13, "type"),
    ("التاريخ",             12, "type"),
    ("الكود / الباركود",    15, "type"),
    ("الصنف كما هو مكتوب",  34, "type"),
    ("الكمية",              9,  "type"),
    ("سعر الفاتورة",        12, "type"),
    ("اسم الصنف في أودو",   34, "auto"),
    ("سعر أودو",            11, "auto"),
    ("الحالة",              26, "auto"),
    ("ملاحظة",              24, "note"),
]
HELP = {
    0: "رقم الفاتورة اليدوية كما هو على الورقة. كل رقم = طلبية مستقلة في أودو.",
    1: "تاريخ الفاتورة الورقية. أودو ستؤرّخ الطلبية بنفس اليوم.",
    2: "اختياري بس الأدق. انسخه من ورقة «الكتالوج» أو امسح الباركود.",
    3: "اكتب الاسم كما هو بالفاتورة. إن تركت الكود فارغاً نطابقه بالاسم.",
    4: "الكمية المباعة.",
    5: "السعر الذي بِيع به فعلاً على الورقة — وليس سعر الكتالوج.",
    6: "يتعبّى تلقائياً — لا تكتب فيه.",
    7: "يتعبّى تلقائياً — سعر الكتالوج للمقارنة.",
    8: "يتعبّى تلقائياً — اقرأه قبل أن ترسل الملف.",
    9: "أي ملاحظة: مرتجع، صنف جديد، اسم المورد…",
}


def fetch(project):
    c = OdooClient(OdooConfig.from_env())
    p = c.search_read("res.partner", ["|", ("name", "=", project), ("name", "ilike", project)],
                      ["name"], limit=1)
    if not p:
        raise SystemExit(f"ما لقينا شريك اسمه «{project}»")
    total = c.search_count("product.template", [])
    rows, off = [], 0
    while off < total:
        page = c.search_read("product.template", [],
                             ["default_code", "name", "list_price", "categ_id"],
                             limit=500, offset=off, order="default_code")
        if not page:
            break
        rows += page
        off += len(page)
    return p[0], rows


def head(ws, labels, widths=None, row=1):
    for i, text in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=i, value=text)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEAD, CENTER, BOX
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 28


def sheet_entry(wb, last):
    ws = wb.create_sheet("فواتير المشروع")
    ws.sheet_view.rightToLeft = True
    head(ws, [c[0] for c in COLS], [c[1] for c in COLS])
    for i, note in HELP.items():
        ws.cell(row=1, column=i + 1).comment = Comment(note, "العون", height=110, width=230)

    code, name, price = f"'الكتالوج'!$A$2:$A${last}", f"'الكتالوج'!$B$2:$B${last}", f"'الكتالوج'!$C$2:$C${last}"
    for r in range(2, 2 + ROWS):
        # code wins when given; otherwise fall back to an exact name match
        ws.cell(row=r, column=7, value=(
            f'=IF($C{r}<>"",IFERROR(INDEX({name},MATCH($C{r},{code},0)),"✖ كود غير موجود"),'
            f'IF($D{r}<>"",IFERROR(INDEX({name},MATCH($D{r},{name},0)),"؟ بحاجة مطابقة"),""))'))
        ws.cell(row=r, column=8, value=(
            f'=IF($G{r}="","",IF(OR(LEFT($G{r},1)="✖",LEFT($G{r},1)="؟"),"",'
            f'IF($C{r}<>"",INDEX({price},MATCH($C{r},{code},0)),INDEX({price},MATCH($D{r},{name},0)))))'))
        ws.cell(row=r, column=9, value=(
            f'=IF($G{r}="","",IF(LEFT($G{r},1)="✖","كود غير موجود — صحّحه",'
            f'IF(LEFT($G{r},1)="؟","صنف جديد أو اسم مختلف",'
            f'IF($H{r}=0,"موجود — بلا سعر بالكتالوج","جاهز ✔"))))'))
        for col in range(1, len(COLS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = BOX
            kind = COLS[col - 1][2]
            if kind == "type":
                cell.font, cell.fill = F_BODY, FILL_TYPE
            elif kind == "auto":
                cell.font, cell.fill = F_AUTO, FILL_AUTO
            else:
                cell.font, cell.fill = F_BODY, FILL_NOTE
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=5).number_format = "#,##0.##"
        ws.cell(row=r, column=6).number_format = "#,##0.000"
        ws.cell(row=r, column=8).number_format = "#,##0.000"
    ws.freeze_panes = "A2"
    return ws


def sheet_catalog(wb, prods):
    ws = wb.create_sheet("الكتالوج")
    ws.sheet_view.rightToLeft = True
    head(ws, ["الكود", "اسم الصنف", "سعر الكتالوج", "الفئة"], [15, 42, 13, 18])
    for r, p in enumerate(prods, start=2):
        ws.cell(row=r, column=1, value=p["default_code"] or "").font = F_SMALL
        ws.cell(row=r, column=2, value=p["name"]).font = F_SMALL
        c3 = ws.cell(row=r, column=3, value=p["list_price"])
        c3.font, c3.number_format = F_SMALL, "#,##0.000"
        ws.cell(row=r, column=4,
                value=p["categ_id"][1] if p["categ_id"] else "—").font = F_SMALL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{len(prods) + 1}"
    return ws


def sheet_help(wb, partner, prods, last):
    ws = wb.create_sheet("التعليمات")
    ws.sheet_view.rightToLeft = True
    for col, w in (("A", 4), ("B", 26), ("C", 82)):
        ws.column_dimensions[col].width = w

    def line(r, a, b, c, font=F_BODY, fill=None):
        for col, val in ((1, a), (2, b), (3, c)):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = font
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill

    zero = sum(1 for p in prods if not p["list_price"])
    line(1, "", "فواتير المشروع اليدوية", partner["name"],
         Font(name="Arial", bold=True, size=14, color=INK))
    line(2, "", "", f"الكتالوج: {len(prods)} صنف · منها {zero} بسعر صفر", F_AUTO)

    rows = [
        ("", "الخطوات", "", F_BOLD, FILL_AUTO),
        ("١", "افتح «فواتير المشروع»", "انقل كل سطر من الفواتير الورقية. الأعمدة الصفراء هي التي تكتب فيها.", F_BODY, None),
        ("٢", "رقم الفاتورة والتاريخ", "كرّرهما في كل سطر من نفس الفاتورة. كل رقم فاتورة يصير طلبية مستقلة بتاريخها.", F_BODY, None),
        ("٣", "الكود أفضل من الاسم", "إن عرفت كود الصنف اكتبه — المطابقة تصير مؤكدة. ابحث عنه في ورقة «الكتالوج».", F_BODY, None),
        ("٤", "سعر الفاتورة", "اكتب السعر الذي بِعت به فعلاً على الورقة. عمود «سعر أودو» للمقارنة فقط.", F_BODY, None),
        ("٥", "راقب عمود الحالة", "لا ترسل الملف قبل أن تقرأه — يخبرك أي سطر جاهز وأيها يحتاج قراراً.", F_BODY, None),
        ("", "", "", F_BODY, None),
        ("", "معاني الحالة", "", F_BOLD, FILL_AUTO),
        ("", "جاهز ✔", "الصنف موجود وله سعر. لا إجراء.", F_BODY, None),
        ("", "موجود — بلا سعر بالكتالوج", "الصنف معرّف لكن سعره صفر. سعر الفاتورة الذي كتبته سيُستعمل، ويمكننا تثبيته في الكتالوج بعدها.", F_BODY, None),
        ("", "صنف جديد أو اسم مختلف", "لم نجد الاسم حرفياً. إمّا صنف جديد ننشئه، أو نفس الصنف باسم مختلف. سأطابقه لك آلياً وأعرض عليك الاحتمالات قبل أي إنشاء.", F_BODY, None),
        ("", "كود غير موجود — صحّحه", "الكود المكتوب ليس في الكتالوج. راجعه.", F_BODY, None),
        ("", "", "", F_BODY, None),
        ("", "ماذا يحدث بعد أن ترسله", "", F_BOLD, FILL_AUTO),
        ("١", "تقرير مطابقة", "أعيد لك قائمة: المطابق، والمتشابه مع الاحتمالات، والجديد فعلاً — قبل أن أنشئ أي صنف.", F_BODY, None),
        ("٢", "الإدخال", "أنشئ الأصناف الجديدة التي توافق عليها، ثم طلبية لكل فاتورة بتاريخها، وأسلّمها وأفوترها.", F_BODY, None),
        ("٣", "مراجعة الأسعار", "أعطيك ملفاً بكل بنود المشروع وسعر كل بند لتراجعها وتضع سعر المقاول، ثم أطبّقها.", F_BODY, None),
        ("٤", "الكشف النهائي", "كشف حساب واحد للمشروع: كل البضاعة، المرتجعات، الدفعات، والرصيد.", F_BODY, None),
    ]
    r = 4
    for a, b, cc, font, fill in rows:
        line(r, a, b, cc, font, fill)
        r += 1

    line(r + 1, "", "مثال على سطر", "", F_BOLD, FILL_AUTO)
    head(ws, [c[0] for c in COLS], None, row=r + 2)
    example = ["INV-114", "2026-07-12", "500086", '3/4" ماسورة (خارجي) سعودي', "3", "6.000",
               "(تلقائي)", "(تلقائي)", "(تلقائي)", "خصم للمقاول"]
    for i, val in enumerate(example, start=1):
        cell = ws.cell(row=r + 3, column=i, value=val)
        cell.font, cell.border = F_BODY, BOX
        cell.fill = FILL_TYPE if COLS[i - 1][2] == "type" else FILL_AUTO
    return ws


def main():
    project = sys.argv[1] if len(sys.argv) > 1 else "مشروع المدرسة"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(f"فواتير_{project}.xlsx")
    partner, prods = fetch(project)
    last = len(prods) + 1
    wb = Workbook()
    wb.remove(wb.active)
    sheet_help(wb, partner, prods, last)
    sheet_entry(wb, last)
    sheet_catalog(wb, prods)
    wb.active = 1
    wb.save(out)
    print(f"المشروع : {partner['name']}")
    print(f"الكتالوج: {len(prods)} صنف (بسعر صفر: {sum(1 for p in prods if not p['list_price'])})")
    print(f"أسطر الإدخال: {ROWS}")
    print(f"✔ {out}")


if __name__ == "__main__":
    main()
